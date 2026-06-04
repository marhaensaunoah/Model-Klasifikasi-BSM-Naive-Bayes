# ============================================================
# SISTEM KLASIFIKASI PENERIMA BSM - NAIVE BAYES
# Dashboard Machine Learning Modern Premium | Streamlit App
# Jalankan: streamlit run app.py
# ============================================================

import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
from sklearn.preprocessing import StandardScaler, LabelEncoder
from sklearn.model_selection import train_test_split
from sklearn.naive_bayes import GaussianNB
from sklearn.metrics import accuracy_score, confusion_matrix, classification_report
import io
import time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# Import CSS dan styling dari file terpisah
from styles import load_custom_css, make_plotly_theme, make_plotly_colors

# ─────────────────────────────────────────────────────────────
# KONFIGURASI HALAMAN AWAL
# ─────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Klasifikasi BSM | Naive Bayes",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ─────────────────────────────────────────────────────────────
# LOAD CUSTOM CSS
# ─────────────────────────────────────────────────────────────
load_custom_css()

# ─────────────────────────────────────────────────────────────
# SESSION STATE INIT
# ─────────────────────────────────────────────────────────────
if 'app_state' not in st.session_state:
    st.session_state.app_state = 'home'
if 'df_raw' not in st.session_state:
    st.session_state.df_raw = None
if 'df_original' not in st.session_state:
    st.session_state.df_original = None
if 'df_processed' not in st.session_state:
    st.session_state.df_processed = None
if 'preprocess_info' not in st.session_state:
    st.session_state.preprocess_info = None
if 'train_results' not in st.session_state:
    st.session_state.train_results = None
if 'test_size' not in st.session_state:
    st.session_state.test_size = 0.2
if 'training_locked' not in st.session_state:
    st.session_state.training_locked = False
if 'batch_prediction_df' not in st.session_state:
    st.session_state.batch_prediction_df = None

# Data Change Log
if 'data_change_log' not in st.session_state:
    st.session_state.data_change_log = {
        'initial_count': 0,
        'current_count': 0,
        'added_count': 0,
        'deleted_count': 0,
        'edited_count': 0,
        'last_activity': None,
        'activities': []
    }

# ─────────────────────────────────────────────────────────────
# HELPER FUNCTIONS (UI)
# ─────────────────────────────────────────────────────────────

def render_sticky_header(icon: str, title: str):
    """Render fixed header yang selalu terlihat di atas saat scroll."""
    st.markdown(f"""
    <div class="sticky-header">
        <div class="sticky-header-icon">{icon}</div>
        <div class="sticky-header-title">{title}</div>
    </div>
    <div class="sticky-header-spacer"></div>
    """, unsafe_allow_html=True)

def render_metric_card(icon: str, label: str, value: str, desc: str = ""):
    """Render single modern metric card."""
    st.markdown(f"""
    <div class="modern-card">
        <div class="metric-icon">{icon}</div>
        <div class="metric-label">{label}</div>
        <div class="metric-value">{value}</div>
        <div class="metric-desc">{desc}</div>
    </div>
    """, unsafe_allow_html=True)

def render_metric_cards(metrics: list, cols_per_row: int = 4):
    """Render multiple metric cards in a row."""
    cols = st.columns(cols_per_row)
    for i, m in enumerate(metrics):
        with cols[i % cols_per_row]:
            render_metric_card(m['icon'], m['label'], m['value'], m.get('desc', ''))

def alert(msg: str, kind: str = "success"):
    """Render alert box."""
    icons = {"success": "✅", "info": "ℹ️", "warning": "⚠️", "error": "❌"}
    st.markdown(f'<div class="alert-{kind}">{icons.get(kind, "•")} {msg}</div>', unsafe_allow_html=True)

def update_change_log(activity_type: str, details: str, count_change: int = 0):
    """Update data change log."""
    log = st.session_state.data_change_log
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    log['activities'].insert(0, {
        'time': now,
        'type': activity_type,
        'details': details,
        'count_change': count_change
    })
    
    if len(log['activities']) > 50:
        log['activities'] = log['activities'][:50]
    
    log['last_activity'] = f"{activity_type}: {details} ({now})"
    
    if activity_type == 'add':
        log['added_count'] += count_change
    elif activity_type == 'delete':
        log['deleted_count'] += count_change
    elif activity_type == 'edit':
        log['edited_count'] += count_change

# ─────────────────────────────────────────────────────────────
# TEMPLATE DATASET
# ─────────────────────────────────────────────────────────────

def create_dataset_template():
    """Membuat template dataset Excel."""
    template_data = {
        'NO': [1, 2],
        'NAMA PESERTA DIDIK': ['Contoh Siswa 1', 'Contoh Siswa 2'],
        'SEKOLAH': ['SMP Negeri 1', 'SMP Negeri 2'],
        'KELAS': ['7', '8'],
        'NIK': ['1234567890123456', '2345678901234567'],
        'NISN': ['1234567890', '2345678901'],
        'NAMA AYAH/IBU': ['Ayah Contoh 1', 'Ibu Contoh 2'],
        'KECAMATAN': ['Kecamatan A', 'Kecamatan B'],
        'BESARAN BIAYA': [1000000, 1200000],
        'PENDAPATAN ORANG TUA': [2000000, 1500000],
        'PEKERJAAN ORANG TUA': ['Buruh', 'Petani'],
        'JUMLAH TANGGUNGAN': [3, 4],
        'STATUS RUMAH': ['Milik Sendiri', 'Kontrak/sewa'],
        'LABEL': ['Ya', 'Tidak']
    }
    return pd.DataFrame(template_data)

def get_template_download():
    """Menyediakan file template untuk download."""
    df_template = create_dataset_template()
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_template.to_excel(writer, index=False, sheet_name='Template BSM')
    output.seek(0)
    return output.getvalue()

# ─────────────────────────────────────────────────────────────
# PREPROCESSING ENGINE (DIPERBAIKI - NO DATA LEAKAGE)
# ─────────────────────────────────────────────────────────────

def fit_preprocessing_from_train(df_train: pd.DataFrame):
    """
    Fit preprocessing (scaler, encoder, imputation values) hanya dari data train.
    """
    info = {}
    
    # Hitung nilai untuk imputasi dari data train
    info['mean_pendapatan'] = pd.to_numeric(df_train['PENDAPATAN ORANG TUA'], errors='coerce').mean()
    info['median_tanggungan'] = pd.to_numeric(df_train['JUMLAH TANGGUNGAN'], errors='coerce').median()
    
    # Fit LabelEncoder untuk pekerjaan
    le_pekerjaan = LabelEncoder()
    le_pekerjaan.fit(df_train['PEKERJAAN ORANG TUA'].astype(str))
    info['le_pekerjaan'] = le_pekerjaan
    
    # Mapping untuk status rumah
    info['status_rumah_mapping'] = {"Milik Sendiri": 1, "Kontrak/sewa": 0}
    
    # Fit StandardScaler untuk pendapatan
    scaler = StandardScaler()
    pendapatan_clean = pd.to_numeric(df_train['PENDAPATAN ORANG TUA'], errors='coerce').fillna(info['mean_pendapatan'])
    scaler.fit(pendapatan_clean.values.reshape(-1, 1))
    info['scaler'] = scaler
    
    # Mapping untuk label
    info['label_mapping'] = {"Ya": 1, "Tidak": 0}
    
    return info

def transform_preprocessing(df: pd.DataFrame, info: dict, is_train: bool = False):
    """
    Transform preprocessing menggunakan info yang sudah di-fit dari data train.
    """
    df = df.copy()
    
    # Konversi numerik
    df['PENDAPATAN ORANG TUA'] = pd.to_numeric(df['PENDAPATAN ORANG TUA'], errors='coerce')
    df['JUMLAH TANGGUNGAN'] = pd.to_numeric(df['JUMLAH TANGGUNGAN'], errors='coerce')
    
    # Imputasi missing values (gunakan nilai dari train)
    df['PENDAPATAN ORANG TUA'] = df['PENDAPATAN ORANG TUA'].fillna(info['mean_pendapatan'])
    df['JUMLAH TANGGUNGAN'] = df['JUMLAH TANGGUNGAN'].fillna(info['median_tanggungan'])
    
    # Ekstrak kelas numerik
    df['KELAS_NUM'] = df['KELAS'].astype(str).str.extract(r'(\d+)')
    df['KELAS_NUM'] = pd.to_numeric(df['KELAS_NUM'], errors='coerce').fillna(0).astype(int)
    
    # Encoding pekerjaan
    le = info['le_pekerjaan']
    # Handle unknown categories
    df['PEKERJAAN ORANG TUA_ENC'] = df['PEKERJAAN ORANG TUA'].astype(str).apply(
        lambda x: le.transform([x])[0] if x in le.classes_ else -1
    )
    
    # Encoding status rumah
    df['STATUS RUMAH_ENC'] = df['STATUS RUMAH'].map(info['status_rumah_mapping'])
    df['STATUS RUMAH_ENC'] = df['STATUS RUMAH_ENC'].fillna(0).astype(int)
    
    # Standardisasi pendapatan (Z-Score)
    scaler = info['scaler']
    df['PENDAPATAN_ZSCORE'] = scaler.transform(df[['PENDAPATAN ORANG TUA']])
    
    # Encoding label
    df['LABEL'] = df['LABEL'].astype(str).str.strip().str.capitalize()
    df['LABEL_ENC'] = df['LABEL'].map(info['label_mapping'])
    
    return df

def run_preprocessing_correct(df_raw: pd.DataFrame, test_size: float = 0.2):
    """
    Pipeline preprocessing yang benar: Split dulu, baru preprocessing.
    """
    df = df_raw.copy()
    
    # Standardisasi nama kolom
    expected_cols = [
        'NO', 'NAMA PESERTA DIDIK', 'SEKOLAH', 'KELAS',
        'NIK', 'NISN', 'NAMA AYAH/IBU', 'KECAMATAN',
        'BESARAN BIAYA', 'PENDAPATAN ORANG TUA',
        'PEKERJAAN ORANG TUA', 'JUMLAH TANGGUNGAN',
        'STATUS RUMAH', 'LABEL'
    ]
    if len(df.columns) == len(expected_cols):
        df.columns = expected_cols
    
    # Hapus duplikat
    before_dup = len(df)
    df = df.drop_duplicates()
    after_dup = len(df)
    dup_removed = before_dup - after_dup
    
    # Validasi label
    df_valid = df[df['LABEL'].notna()].copy()
    df_valid['LABEL'] = df_valid['LABEL'].astype(str).str.strip().str.capitalize()
    df_valid = df_valid[df_valid['LABEL'].isin(['Ya', 'Tidak'])].copy()
    
    # Hitung missing values awal
    missing_before = df_valid.isnull().sum().to_dict()
    
    # Split data (stratify berdasarkan label)
    y = df_valid['LABEL']
    if len(y.unique()) >= 2 and test_size > 0 and test_size < 1:
        train_idx, test_idx = train_test_split(
            df_valid.index, test_size=test_size, random_state=42, stratify=y
        )
        df_train_raw = df_valid.loc[train_idx].copy()
        df_test_raw = df_valid.loc[test_idx].copy()
    elif test_size == 0.0:
        df_train_raw = df_valid.copy()
        df_test_raw = pd.DataFrame(columns=df_valid.columns)
    else:  # test_size == 1.0
        df_train_raw = pd.DataFrame(columns=df_valid.columns)
        df_test_raw = df_valid.copy()
    
    # Fit preprocessing hanya dari data train
    if len(df_train_raw) > 0:
        preprocess_info = fit_preprocessing_from_train(df_train_raw)
    else:
        # Jika tidak ada data train, gunakan dummy info
        preprocess_info = {
            'mean_pendapatan': 0,
            'median_tanggungan': 0,
            'le_pekerjaan': LabelEncoder(),
            'status_rumah_mapping': {"Milik Sendiri": 1, "Kontrak/sewa": 0},
            'scaler': StandardScaler(),
            'label_mapping': {"Ya": 1, "Tidak": 0}
        }
        preprocess_info['le_pekerjaan'].fit(['Buruh', 'Petani', 'Wiraswasta', 'PNS', 'Lainnya'])
    
    # Transform data train dan test
    if len(df_train_raw) > 0:
        df_train = transform_preprocessing(df_train_raw, preprocess_info, is_train=True)
    else:
        df_train = pd.DataFrame()
    
    if len(df_test_raw) > 0:
        df_test = transform_preprocessing(df_test_raw, preprocess_info, is_train=False)
    else:
        df_test = pd.DataFrame()
    
    # Hitung missing values akhir
    missing_after_train = df_train.isnull().sum().to_dict() if len(df_train) > 0 else {}
    missing_after_test = df_test.isnull().sum().to_dict() if len(df_test) > 0 else {}
    
    # Gabungkan kembali untuk ditampilkan (dengan informasi asal)
    if len(df_train) > 0:
        df_train['_DATASET_TYPE'] = 'TRAIN'
    if len(df_test) > 0:
        df_test['_DATASET_TYPE'] = 'TEST'
    
    if len(df_test) > 0 and len(df_train) > 0:
        df_combined = pd.concat([df_train, df_test], ignore_index=True)
    elif len(df_train) > 0:
        df_combined = df_train
    elif len(df_test) > 0:
        df_combined = df_test
    else:
        df_combined = pd.DataFrame()
    
    info = {
        'preprocess_info': preprocess_info,
        'df_train': df_train,
        'df_test': df_test,
        'missing_before': missing_before,
        'missing_after_train': missing_after_train,
        'dup_removed': dup_removed,
        'total_rows': len(df_combined),
        'train_rows': len(df_train),
        'test_rows': len(df_test),
        'test_size': test_size
    }
    
    # Cetak info ke terminal
    print_preprocessing_info_correct(df_train, df_test, info)
    
    return df_combined, info

def print_preprocessing_info_correct(df_train, df_test, info):
    """Cetak informasi preprocessing ke terminal."""
    print("\n" + "="*80)
    print("PREPROCESSING DATA (TANPA DATA LEAKAGE)")
    print("="*80)
    
    print(f"\n[1] DATA AWAL")
    print(f"    Jumlah data awal : {info['total_rows'] + info['dup_removed']} baris")
    print(f"    Data duplikat    : {info['dup_removed']} baris")
    
    print(f"\n[2] SPLIT DATA")
    print(f"    Data Training    : {info['train_rows']} baris")
    print(f"    Data Testing     : {info['test_rows']} baris")
    
    print(f"\n[3] PREPROCESSING FIT (dari data TRAIN saja)")
    pre = info['preprocess_info']
    print(f"    Mean pendapatan (train) : {pre['mean_pendapatan']:.2f}")
    print(f"    Median tanggungan (train): {pre['median_tanggungan']:.2f}")
    print(f"    Scaler mean (train)      : {pre['scaler'].mean_[0]:.6f}")
    print(f"    Scaler std (train)       : {pre['scaler'].scale_[0]:.6f}")
    
    print(f"\n[4] MAPPING ENCODING (dari data TRAIN)")
    le = pre['le_pekerjaan']
    print(f"\n    Pekerjaan Orang Tua → Label Encoding:")
    for kd, pk in zip(le.transform(le.classes_), le.classes_):
        print(f"    {pk:30s} → {kd}")
    
    print("\n" + "="*80)
    print("PREPROCESSING SELESAI")
    print("="*80 + "\n")

# ─────────────────────────────────────────────────────────────
# TRAINING ENGINE (DIPERBAIKI)
# ─────────────────────────────────────────────────────────────

def run_training_correct(df_train: pd.DataFrame, df_test: pd.DataFrame):
    """
    Latih model Gaussian Naive Bayes menggunakan data train yang sudah diproses.
    """
    feature_cols = [
        'KELAS_NUM',
        'PENDAPATAN_ZSCORE',
        'PEKERJAAN ORANG TUA_ENC',
        'JUMLAH TANGGUNGAN',
        'STATUS RUMAH_ENC'
    ]
    
    X_train = df_train[feature_cols].copy()
    y_train = df_train['LABEL_ENC'].copy()
    
    # Validasi
    valid_train = X_train.notna().all(axis=1) & y_train.notna()
    X_train = X_train[valid_train]
    y_train = y_train[valid_train]
    
    print("\n" + "="*80)
    print("TRAINING GAUSSIAN NAIVE BAYES")
    print("="*80)
    
    print(f"\n[1] DATA TRAINING")
    print(f"    Total data valid : {len(X_train)} baris")
    
    # Cek distribusi kelas
    train_dist = y_train.value_counts()
    print(f"\n    Distribusi Kelas:")
    print(f"    - Tidak Penerima (0) : {train_dist.get(0, 0)} ({train_dist.get(0, 0)/len(y_train)*100:.1f}%)")
    print(f"    - Penerima (1)       : {train_dist.get(1, 0)} ({train_dist.get(1, 0)/len(y_train)*100:.1f}%)")
    
    # Warning untuk imbalanced dataset
    min_class_pct = min(train_dist.get(0, 0), train_dist.get(1, 0)) / len(y_train) * 100
    if min_class_pct < 20:
        print(f"\n    ⚠️ PERINGATAN: Dataset tidak seimbang! Kelas minoritas hanya {min_class_pct:.1f}%")
    
    # Training model
    model = GaussianNB()
    model.fit(X_train, y_train)
    
    print(f"\n[2] PRIOR PROBABILITY")
    print(f"    P(Tidak) = {model.class_prior_[0]:.6f} ({model.class_prior_[0]*100:.2f}%)")
    print(f"    P(Ya)    = {model.class_prior_[1]:.6f} ({model.class_prior_[1]*100:.2f}%)")
    
    print(f"\n[3] MEAN (THETA_) PER FITUR")
    print(f"    {'Fitur':35s} {'Tidak (0)':>12s} {'Ya (1)':>12s}")
    print(f"    {'-'*35} {'-'*12} {'-'*12}")
    for i, feat in enumerate(feature_cols):
        print(f"    {feat:35s} {model.theta_[0][i]:12.6f} {model.theta_[1][i]:12.6f}")
    
    print(f"\n[4] VARIANCE (VAR_) PER FITUR")
    print(f"    {'Fitur':35s} {'Tidak (0)':>12s} {'Ya (1)':>12s}")
    print(f"    {'-'*35} {'-'*12} {'-'*12}")
    for i, feat in enumerate(feature_cols):
        print(f"    {feat:35s} {model.var_[0][i]:12.6f} {model.var_[1][i]:12.6f}")
    
    # Prediksi pada data train
    y_train_pred = model.predict(X_train)
    train_acc = accuracy_score(y_train, y_train_pred)
    print(f"\n[5] AKURASI TRAINING")
    print(f"    Akurasi Training : {train_acc:.6f} ({train_acc*100:.2f}%)")
    
    # Evaluasi pada data test jika ada
    test_acc = None
    y_pred = np.array([])
    y_pred_proba = np.array([])
    classification_dict = None
    cm = np.array([[0, 0], [0, 0]])
    
    if len(df_test) > 0:
        X_test = df_test[feature_cols].copy()
        y_test = df_test['LABEL_ENC'].copy()
        
        valid_test = X_test.notna().all(axis=1) & y_test.notna()
        X_test = X_test[valid_test]
        y_test = y_test[valid_test]
        
        if len(X_test) > 0:
            y_pred = model.predict(X_test)
            y_pred_proba = model.predict_proba(X_test)
            test_acc = accuracy_score(y_test, y_pred)
            cm = confusion_matrix(y_test, y_pred)
            
            # Classification report
            classification_dict = classification_report(y_test, y_pred, 
                                                        target_names=['Tidak Penerima', 'Penerima'],
                                                        output_dict=True)
            
            print(f"\n[6] EVALUASI PADA DATA TESTING")
            print(f"    Total data test valid : {len(X_test)} baris")
            print(f"    Akurasi Testing       : {test_acc:.6f} ({test_acc*100:.2f}%)")
            print(f"\n    Classification Report:")
            print(classification_report(y_test, y_pred, target_names=['Tidak Penerima', 'Penerima']))
    
    print("\n" + "="*80)
    print("TRAINING SELESAI")
    print("="*80 + "\n")
    
    # Hitung metrik dari classification dict
    if classification_dict:
        prec0 = classification_dict['Tidak Penerima']['precision']
        rec0 = classification_dict['Tidak Penerima']['recall']
        f1_0 = classification_dict['Tidak Penerima']['f1-score']
        prec1 = classification_dict['Penerima']['precision']
        rec1 = classification_dict['Penerima']['recall']
        f1_1 = classification_dict['Penerima']['f1-score']
        sup0 = classification_dict['Tidak Penerima']['support']
        sup1 = classification_dict['Penerima']['support']
        macro_prec = classification_dict['macro avg']['precision']
        macro_rec = classification_dict['macro avg']['recall']
        macro_f1 = classification_dict['macro avg']['f1-score']
        w_prec = classification_dict['weighted avg']['precision']
        w_rec = classification_dict['weighted avg']['recall']
        w_f1 = classification_dict['weighted avg']['f1-score']
    else:
        prec0 = rec0 = f1_0 = prec1 = rec1 = f1_1 = sup0 = sup1 = 0
        macro_prec = macro_rec = macro_f1 = w_prec = w_rec = w_f1 = 0
    
    # Buat dataframe hasil prediksi
    results_df = pd.DataFrame()
    if len(y_pred) > 0:
        results_df = pd.DataFrame({
            'No': range(1, len(y_test) + 1),
            'Actual': ['Penerima' if v == 1 else 'Tidak Penerima' for v in y_test.values],
            'Predicted': ['Penerima' if v == 1 else 'Tidak Penerima' for v in y_pred],
            'Prob Tidak (%)': (y_pred_proba[:, 0] * 100).round(2),
            'Prob Ya (%)': (y_pred_proba[:, 1] * 100).round(2),
            'Status': ['✓ Benar' if a == p else '✗ Salah'
                       for a, p in zip(y_test.values, y_pred)]
        })
    
    return {
        'model': model,
        'X_train': X_train, 'y_train': y_train,
        'X_test': X_test if len(df_test) > 0 else pd.DataFrame(),
        'y_test': y_test if len(df_test) > 0 else pd.Series(),
        'y_pred': y_pred,
        'y_pred_proba': y_pred_proba,
        'train_acc': train_acc,
        'test_acc': test_acc,
        'cm': cm,
        'prec0': prec0, 'rec0': rec0, 'f1_0': f1_0,
        'prec1': prec1, 'rec1': rec1, 'f1_1': f1_1,
        'macro_prec': macro_prec, 'macro_rec': macro_rec, 'macro_f1': macro_f1,
        'w_prec': w_prec, 'w_rec': w_rec, 'w_f1': w_f1,
        'sup0': sup0, 'sup1': sup1,
        'feature_cols': feature_cols,
        'classification_report': classification_dict,
        'results_df': results_df,
        'train_distribution': train_dist.to_dict(),
        'class_prior': model.class_prior_,
        'theta': model.theta_,
        'var': model.var_
    }

# ─────────────────────────────────────────────────────────────
# FUNGSI UNTUK EDIT/ADD/DELETE DATA
# ─────────────────────────────────────────────────────────────

def add_new_data_to_dataset(new_row: dict):
    """Menambahkan data baru ke dataset."""
    if st.session_state.df_raw is not None:
        df = st.session_state.df_raw.copy()
        
        # Tentukan NO berikutnya
        if 'NO' in df.columns and len(df) > 0:
            max_no = pd.to_numeric(df['NO'], errors='coerce').max()
            new_no = int(max_no) + 1 if pd.notna(max_no) else len(df) + 1
        else:
            new_no = len(df) + 1
        
        new_row['NO'] = new_no
        
        # Tambahkan ke dataframe
        new_df = pd.DataFrame([new_row])
        df = pd.concat([df, new_df], ignore_index=True)
        
        # Update session state
        st.session_state.df_raw = df
        st.session_state.df_original = df.copy()
        
        # Reset proses yang tergantung
        st.session_state.df_processed = None
        st.session_state.preprocess_info = None
        st.session_state.train_results = None
        st.session_state.training_locked = False
        
        # Update data change log
        current_count = len(df)
        st.session_state.data_change_log['current_count'] = current_count
        
        update_change_log('add', f"Menambahkan data baru (NO: {new_no})", 1)
        
        return True, new_no
    return False, None

def delete_selected_rows(indices_to_delete: list):
    """Menghapus baris yang dipilih dari dataset."""
    if st.session_state.df_raw is not None and indices_to_delete:
        df = st.session_state.df_raw.copy()
        before_count = len(df)
        
        df = df.drop(indices_to_delete).reset_index(drop=True)
        
        # Update NO jika ada
        if 'NO' in df.columns:
            df['NO'] = range(1, len(df) + 1)
        
        after_count = len(df)
        deleted_count = before_count - after_count
        
        st.session_state.df_raw = df
        st.session_state.df_original = df.copy()
        
        # Reset proses yang tergantung
        st.session_state.df_processed = None
        st.session_state.preprocess_info = None
        st.session_state.train_results = None
        st.session_state.training_locked = False
        
        # Update data change log
        st.session_state.data_change_log['current_count'] = after_count
        
        update_change_log('delete', f"Menghapus {deleted_count} baris data", deleted_count)
        
        return True, deleted_count
    return False, 0

def save_edited_data(edited_df: pd.DataFrame):
    """Menyimpan hasil edit data."""
    if st.session_state.df_raw is not None:
        before_count = len(st.session_state.df_raw)
        
        st.session_state.df_raw = edited_df.copy()
        st.session_state.df_original = edited_df.copy()
        
        # Reset proses yang tergantung
        st.session_state.df_processed = None
        st.session_state.preprocess_info = None
        st.session_state.train_results = None
        st.session_state.training_locked = False
        
        # Update data change log
        st.session_state.data_change_log['current_count'] = len(edited_df)
        
        update_change_log('edit', f"Menyimpan {len(edited_df)} baris data yang diedit", 0)
        
        return True
    return False

def reset_to_original():
    """Reset ke data awal."""
    if st.session_state.df_original is not None:
        st.session_state.df_raw = st.session_state.df_original.copy()
        st.session_state.df_processed = None
        st.session_state.preprocess_info = None
        st.session_state.train_results = None
        st.session_state.training_locked = False
        
        # Update data change log
        st.session_state.data_change_log['current_count'] = len(st.session_state.df_original)
        
        update_change_log('reset', "Reset ke data awal", 0)
        
        return True
    return False

# ─────────────────────────────────────────────────────────────
# EXCEL EXPORT HELPERS
# ─────────────────────────────────────────────────────────────

def _safe_excel_value(val):
    """Konversi nilai pandas NA menjadi None untuk kompatibilitas Excel."""
    if pd.isna(val):
        return None
    return val

def _style_header(cell, bg="065f46", fg="FFFFFF"):
    cell.font = Font(name='Inter', bold=True, color=fg, size=11)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='BFBFBF')
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

def _style_data(cell, bg="FFFFFF", bold=False, align='center'):
    cell.font = Font(name='Inter', size=10, bold=bold)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal=align, vertical='center')
    thin = Side(style='thin', color='D9D9D9')
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

def export_model_summary_excel(train_results: dict, preprocess_info: dict) -> bytes:
    """Export ringkasan model ke Excel."""
    wb = Workbook()
    
    # Sheet 1: Ringkasan Model
    ws1 = wb.active
    ws1.title = "Ringkasan Model"
    
    ws1.merge_cells('A1:D1')
    ws1['A1'].value = "RINGKASAN MODEL KLASIFIKASI BSM"
    ws1['A1'].font = Font(name='Inter', bold=True, size=14, color='FFFFFF')
    ws1['A1'].fill = PatternFill("solid", start_color="065f46")
    ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[1].height = 30
    
    # Informasi dataset
    row = 3
    _style_header(ws1.cell(row=row, column=1, value="Parameter"), "059669")
    _style_header(ws1.cell(row=row, column=2, value="Nilai"), "059669")
    _style_header(ws1.cell(row=row, column=3, value="Keterangan"), "059669")
    
    row += 1
    _style_data(ws1.cell(row=row, column=1, value="Jumlah Data Training"), "FFFFFF", align='left')
    _style_data(ws1.cell(row=row, column=2, value=str(len(train_results['X_train']))), "FFFFFF")
    _style_data(ws1.cell(row=row, column=3, value="Data yang digunakan untuk melatih model"), "FFFFFF", align='left')
    
    row += 1
    _style_data(ws1.cell(row=row, column=1, value="Jumlah Data Testing"), "ecfdf5", align='left')
    _style_data(ws1.cell(row=row, column=2, value=str(len(train_results['X_test']))), "ecfdf5")
    _style_data(ws1.cell(row=row, column=3, value="Data yang digunakan untuk evaluasi"), "ecfdf5", align='left')
    
    row += 1
    _style_data(ws1.cell(row=row, column=1, value="Jumlah Fitur"), "FFFFFF", align='left')
    _style_data(ws1.cell(row=row, column=2, value=str(len(train_results['feature_cols']))), "FFFFFF")
    _style_data(ws1.cell(row=row, column=3, value="Fitur input yang digunakan"), "FFFFFF", align='left')
    
    row += 2
    _style_header(ws1.cell(row=row, column=1, value="Daftar Fitur"), "059669")
    _style_header(ws1.cell(row=row, column=2, value="Tipe"), "059669")
    _style_header(ws1.cell(row=row, column=3, value="Deskripsi"), "059669")
    
    fitur_desc = {
        'KELAS_NUM': 'Numerik', 'PENDAPATAN_ZSCORE': 'Numerik (Standardized)',
        'PEKERJAAN ORANG TUA_ENC': 'Kategorikal (Encoded)', 'JUMLAH TANGGUNGAN': 'Numerik',
        'STATUS RUMAH_ENC': 'Kategorikal (Binary)'
    }
    
    for i, feat in enumerate(train_results['feature_cols']):
        row += 1
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(ws1.cell(row=row, column=1, value=feat), bg, align='left')
        _style_data(ws1.cell(row=row, column=2, value=fitur_desc.get(feat, 'Numerik')), bg)
        _style_data(ws1.cell(row=row, column=3, value="Fitur untuk klasifikasi BSM"), bg, align='left')
    
    # Sheet 2: Prior Probability & Parameter
    ws2 = wb.create_sheet("Prior & Parameter")
    
    ws2.merge_cells('A1:B1')
    ws2['A1'].value = "PRIOR PROBABILITY"
    ws2['A1'].font = Font(name='Inter', bold=True, size=12, color='FFFFFF')
    ws2['A1'].fill = PatternFill("solid", start_color="065f46")
    ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    row = 3
    _style_header(ws2.cell(row=row, column=1, value="Kelas"), "059669")
    _style_header(ws2.cell(row=row, column=2, value="Prior Probability"), "059669")
    _style_header(ws2.cell(row=row, column=3, value="Persentase"), "059669")
    
    if train_results['class_prior'] is not None:
        row += 1
        _style_data(ws2.cell(row=row, column=1, value="Tidak Penerima (0)"), "FFFFFF", align='left')
        _style_data(ws2.cell(row=row, column=2, value=f"{train_results['class_prior'][0]:.6f}"), "FFFFFF")
        _style_data(ws2.cell(row=row, column=3, value=f"{train_results['class_prior'][0]*100:.2f}%"), "FFFFFF")
        
        row += 1
        _style_data(ws2.cell(row=row, column=1, value="Penerima (1)"), "ecfdf5", align='left')
        _style_data(ws2.cell(row=row, column=2, value=f"{train_results['class_prior'][1]:.6f}"), "ecfdf5")
        _style_data(ws2.cell(row=row, column=3, value=f"{train_results['class_prior'][1]*100:.2f}%"), "ecfdf5")
    
    row += 2
    ws2.merge_cells(f'A{row}:D{row}')
    _style_header(ws2.cell(row=row, column=1, value="MEAN (THETA) PER FITUR"), "047857")
    
    row += 1
    _style_header(ws2.cell(row=row, column=1, value="Fitur"), "059669")
    _style_header(ws2.cell(row=row, column=2, value="Mean - Tidak Penerima"), "059669")
    _style_header(ws2.cell(row=row, column=3, value="Mean - Penerima"), "059669")
    
    if train_results['theta'] is not None:
        for i, feat in enumerate(train_results['feature_cols']):
            row += 1
            bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
            _style_data(ws2.cell(row=row, column=1, value=feat), bg, align='left')
            _style_data(ws2.cell(row=row, column=2, value=round(train_results['theta'][0][i], 6)), bg)
            _style_data(ws2.cell(row=row, column=3, value=round(train_results['theta'][1][i], 6)), bg)
    
    row += 2
    ws2.merge_cells(f'A{row}:D{row}')
    _style_header(ws2.cell(row=row, column=1, value="VARIANCE (VAR) PER FITUR"), "047857")
    
    row += 1
    _style_header(ws2.cell(row=row, column=1, value="Fitur"), "059669")
    _style_header(ws2.cell(row=row, column=2, value="Variance - Tidak Penerima"), "059669")
    _style_header(ws2.cell(row=row, column=3, value="Variance - Penerima"), "059669")
    
    if train_results['var'] is not None:
        for i, feat in enumerate(train_results['feature_cols']):
            row += 1
            bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
            _style_data(ws2.cell(row=row, column=1, value=feat), bg, align='left')
            _style_data(ws2.cell(row=row, column=2, value=round(train_results['var'][0][i], 6)), bg)
            _style_data(ws2.cell(row=row, column=3, value=round(train_results['var'][1][i], 6)), bg)
    
    # Sheet 3: Evaluasi Model
    ws3 = wb.create_sheet("Evaluasi Model")
    
    ws3.merge_cells('A1:C1')
    ws3['A1'].value = "METRIK EVALUASI MODEL"
    ws3['A1'].font = Font(name='Inter', bold=True, size=12, color='FFFFFF')
    ws3['A1'].fill = PatternFill("solid", start_color="065f46")
    ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    row = 3
    _style_header(ws3.cell(row=row, column=1, value="Metrik"), "059669")
    _style_header(ws3.cell(row=row, column=2, value="Nilai"), "059669")
    _style_header(ws3.cell(row=row, column=3, value="Persentase"), "059669")
    
    metrics_data = [
        ("Akurasi Training", train_results['train_acc']),
        ("Akurasi Testing", train_results['test_acc']),
        ("Precision - Tidak Penerima", train_results['prec0']),
        ("Recall - Tidak Penerima", train_results['rec0']),
        ("F1-Score - Tidak Penerima", train_results['f1_0']),
        ("Precision - Penerima", train_results['prec1']),
        ("Recall - Penerima", train_results['rec1']),
        ("F1-Score - Penerima", train_results['f1_1']),
        ("Macro Avg Precision", train_results['macro_prec']),
        ("Macro Avg Recall", train_results['macro_rec']),
        ("Macro Avg F1-Score", train_results['macro_f1']),
        ("Weighted Avg Precision", train_results['w_prec']),
        ("Weighted Avg Recall", train_results['w_rec']),
        ("Weighted Avg F1-Score", train_results['w_f1']),
    ]
    
    for i, (name, value) in enumerate(metrics_data):
        row += 1
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(ws3.cell(row=row, column=1, value=name), bg, align='left')
        if value is not None:
            _style_data(ws3.cell(row=row, column=2, value=round(value, 6)), bg)
            _style_data(ws3.cell(row=row, column=3, value=f"{value*100:.2f}%"), bg)
        else:
            _style_data(ws3.cell(row=row, column=2, value="N/A"), bg)
            _style_data(ws3.cell(row=row, column=3, value="N/A"), bg)
    
    for ws in [ws1, ws2, ws3]:
        for col in ['A', 'B', 'C', 'D']:
            ws.column_dimensions[col].width = 28
    
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

def export_preprocessing_excel(df: pd.DataFrame, le_pekerjaan: LabelEncoder) -> bytes:
    """Export data preprocessing ke Excel dengan styling."""
    cols_to_export = [
        'NO', 'NAMA PESERTA DIDIK', 'SEKOLAH', 'KELAS', 'KELAS_NUM',
        'PENDAPATAN ORANG TUA', 'PENDAPATAN_ZSCORE',
        'PEKERJAAN ORANG TUA', 'PEKERJAAN ORANG TUA_ENC',
        'JUMLAH TANGGUNGAN', 'STATUS RUMAH', 'STATUS RUMAH_ENC',
        'LABEL', 'LABEL_ENC'
    ]
    
    # Filter kolom yang tersedia
    available_cols = [c for c in cols_to_export if c in df.columns]
    df_exp = df[available_cols].copy()
    
    # Tambahkan tipe dataset jika ada
    if '_DATASET_TYPE' in df.columns:
        df_exp['_DATASET_TYPE'] = df['_DATASET_TYPE']
    
    for col in ['KELAS_NUM', 'PEKERJAAN ORANG TUA_ENC', 'STATUS RUMAH_ENC', 'LABEL_ENC']:
        if col in df_exp.columns:
            df_exp[col] = pd.to_numeric(df_exp[col], errors='coerce')
    
    # Rename columns
    rename_map = {
        'NO': 'No',
        'NAMA PESERTA DIDIK': 'Nama Peserta Didik',
        'SEKOLAH': 'Sekolah',
        'KELAS': 'Kelas',
        'KELAS_NUM': 'Kelas (Angka)',
        'PENDAPATAN ORANG TUA': 'Pendapatan Orang Tua (Rp)',
        'PENDAPATAN_ZSCORE': 'Pendapatan (Z-Score)',
        'PEKERJAAN ORANG TUA': 'Pekerjaan Orang Tua',
        'PEKERJAAN ORANG TUA_ENC': 'Pekerjaan (Encoded)',
        'JUMLAH TANGGUNGAN': 'Jumlah Tanggungan',
        'STATUS RUMAH': 'Status Rumah',
        'STATUS RUMAH_ENC': 'Status Rumah (Encoded)',
        'LABEL': 'Label',
        'LABEL_ENC': 'Label (Encoded)',
        '_DATASET_TYPE': 'Tipe Dataset'
    }
    df_exp = df_exp.rename(columns=rename_map)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Preprocessing"
    
    ws.merge_cells('A1:N1')
    tc = ws['A1']
    tc.value = "DATA HASIL PREPROCESSING - BSM 2022-2024"
    tc.font = Font(name='Inter', bold=True, size=14, color='FFFFFF')
    tc.fill = PatternFill("solid", start_color="065f46")
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30
    
    ws.merge_cells('A2:N2')
    sc = ws['A2']
    sc.value = f"Total Data: {len(df_exp)} baris | Fitur: 5 | Target: LABEL (Ya=1, Tidak=0)"
    sc.font = Font(name='Inter', size=9, italic=True, color='595959')
    sc.fill = PatternFill("solid", start_color="d1fae5")
    sc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 18
    
    headers = list(df_exp.columns)
    for ci, h in enumerate(headers, 1):
        _style_header(ws.cell(row=3, column=ci, value=h))
    ws.row_dimensions[3].height = 32
    
    row_colors = ["FFFFFF", "ecfdf5"]
    for ri, row in enumerate(df_exp.itertuples(index=False), 4):
        bg = row_colors[(ri - 4) % 2]
        for ci, val in enumerate(row, 1):
            safe_val = _safe_excel_value(val)
            cell = ws.cell(row=ri, column=ci, value=safe_val)
            _style_data(cell, bg)
            if ci == 6:  # Pendapatan Orang Tua
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif ci == 7:  # Z-Score
                cell.number_format = '0.000000'
                cell.alignment = Alignment(horizontal='right', vertical='center')
    
    # Set column widths
    default_widths = [5, 28, 28, 10, 12, 22, 18, 22, 16, 16, 18, 18, 8, 14, 12]
    for i, w in enumerate(default_widths[:len(headers)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    
    wm = wb.create_sheet("Mapping Encoding")
    wm.merge_cells('A1:C1')
    wm['A1'].value = "MAPPING ENCODING"
    wm['A1'].font = Font(name='Inter', bold=True, size=13, color='FFFFFF')
    wm['A1'].fill = PatternFill("solid", start_color="065f46")
    wm['A1'].alignment = Alignment(horizontal='center', vertical='center')
    wm.row_dimensions[1].height = 26
    
    _style_header(wm.cell(row=3, column=1, value="Pekerjaan Orang Tua"), "059669")
    _style_header(wm.cell(row=3, column=2, value="Kode"), "059669")
    for i, (kd, pk) in enumerate(zip(le_pekerjaan.transform(le_pekerjaan.classes_), le_pekerjaan.classes_), 4):
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(wm.cell(row=i, column=1, value=pk), bg, align='left')
        _style_data(wm.cell(row=i, column=2, value=int(kd)), bg)
    
    rs = 4 + len(le_pekerjaan.classes_) + 2
    _style_header(wm.cell(row=rs, column=1, value="Status Rumah"), "059669")
    _style_header(wm.cell(row=rs, column=2, value="Kode"), "059669")
    for i, (s, k) in enumerate([("Kontrak/sewa", 0), ("Milik Sendiri", 1)], rs + 1):
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(wm.cell(row=i, column=1, value=s), bg, align='left')
        _style_data(wm.cell(row=i, column=2, value=k), bg)
    
    rs2 = rs + 4
    _style_header(wm.cell(row=rs2, column=1, value="Label"), "059669")
    _style_header(wm.cell(row=rs2, column=2, value="Kode"), "059669")
    for i, (lb, k) in enumerate([("Tidak", 0), ("Ya", 1)], rs2 + 1):
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(wm.cell(row=i, column=1, value=lb), bg, align='left')
        _style_data(wm.cell(row=i, column=2, value=k), bg)
    
    wm.column_dimensions['A'].width = 26
    wm.column_dimensions['B'].width = 10
    
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def export_normalisasi_excel(df: pd.DataFrame, info: dict) -> bytes:
    """Export hasil normalisasi ke Excel."""
    scaler = info['preprocess_info']['scaler']
    mean_val = scaler.mean_[0]
    std_val = scaler.scale_[0]
    
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = "Hasil Normalisasi"
    
    _style_header(ws1.cell(row=1, column=1, value="No"))
    _style_header(ws1.cell(row=1, column=2, value="Nama Peserta Didik"))
    _style_header(ws1.cell(row=1, column=3, value="Pendapatan Asli (Rp)"))
    _style_header(ws1.cell(row=1, column=4, value="Pendapatan Z-Score"))
    _style_header(ws1.cell(row=1, column=5, value="Rumus Z-Score"))
    
    # Filter baris yang memiliki data valid
    valid_rows = df['PENDAPATAN ORANG TUA'].notna()
    df_filtered = df[valid_rows].copy()
    
    for i, (_, row) in enumerate(df_filtered.iterrows(), 2):
        z = (row['PENDAPATAN ORANG TUA'] - mean_val) / std_val
        rumus = f"({row['PENDAPATAN ORANG TUA']} - {mean_val:.2f}) / {std_val:.2f}"
        
        bg = "FFFFFF" if i % 2 == 0 else "ecfdf5"
        _style_data(ws1.cell(row=i, column=1, value=_safe_excel_value(row.get('NO', i-1))), bg)
        _style_data(ws1.cell(row=i, column=2, value=_safe_excel_value(row.get('NAMA PESERTA DIDIK', '-'))), bg, align='left')
        _style_data(ws1.cell(row=i, column=3, value=_safe_excel_value(row['PENDAPATAN ORANG TUA'])), bg)
        ws1.cell(row=i, column=3).number_format = '#,##0'
        _style_data(ws1.cell(row=i, column=4, value=round(z, 6)), bg)
        _style_data(ws1.cell(row=i, column=5, value=rumus), bg, align='left')
    
    ws1.column_dimensions['A'].width = 6
    ws1.column_dimensions['B'].width = 30
    ws1.column_dimensions['C'].width = 22
    ws1.column_dimensions['D'].width = 18
    ws1.column_dimensions['E'].width = 35
    
    ws2 = wb.create_sheet("Parameter Normalisasi")
    _style_header(ws2.cell(row=1, column=1, value="Parameter"), "059669")
    _style_header(ws2.cell(row=1, column=2, value="Nilai"), "059669")
    _style_header(ws2.cell(row=1, column=3, value="Keterangan"), "059669")
    
    _style_data(ws2.cell(row=2, column=1, value="Mean (μ)"), "FFFFFF", align='left')
    _style_data(ws2.cell(row=2, column=2, value=round(mean_val, 4)), "FFFFFF")
    _style_data(ws2.cell(row=2, column=3, value="Rata-rata pendapatan orang tua (dari data train)"), "FFFFFF", align='left')
    
    _style_data(ws2.cell(row=3, column=1, value="Standard Deviation (σ)"), "ecfdf5", align='left')
    _style_data(ws2.cell(row=3, column=2, value=round(std_val, 4)), "ecfdf5")
    _style_data(ws2.cell(row=3, column=3, value="Simpangan baku pendapatan orang tua (dari data train)"), "ecfdf5", align='left')
    
    ws2.merge_cells('A5:C5')
    _style_header(ws2.cell(row=5, column=1, value="Rumus Z-Score: z = (x - μ) / σ"), "047857")
    
    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 16
    ws2.column_dimensions['C'].width = 35
    
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def export_batch_prediction_excel(df_result: pd.DataFrame) -> bytes:
    """Export hasil prediksi batch ke Excel dengan styling profesional."""
    wb = Workbook()
    
    ws = wb.active
    ws.title = "Hasil Prediksi"
    
    ws.merge_cells('A1:H1')
    tc = ws['A1']
    tc.value = "HASIL PREDIKSI PENERIMA BSM - NAIVE BAYES"
    tc.font = Font(name='Inter', bold=True, size=14, color='FFFFFF')
    tc.fill = PatternFill("solid", start_color="065f46")
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30
    
    ws.merge_cells('A2:H2')
    sc = ws['A2']
    sc.value = f"Total Data: {len(df_result)} | Model: Gaussian Naive Bayes | Tanggal: {pd.Timestamp.now().strftime('%d-%m-%Y %H:%M')}"
    sc.font = Font(name='Inter', size=9, italic=True, color='595959')
    sc.fill = PatternFill("solid", start_color="d1fae5")
    sc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 18
    
    headers = list(df_result.columns)
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        _style_header(cell)
    ws.row_dimensions[3].height = 32
    
    row_colors = ["FFFFFF", "ecfdf5"]
    for ri, (_, row) in enumerate(df_result.iterrows(), 4):
        bg = row_colors[(ri - 4) % 2]
        for ci, val in enumerate(row, 1):
            safe_val = _safe_excel_value(val)
            cell = ws.cell(row=ri, column=ci, value=safe_val)
            
            col_name = headers[ci-1] if ci-1 < len(headers) else ""
            if col_name in ['PROB_TIDAK (%)', 'PROB_YA (%)']:
                cell.number_format = '0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
                _style_data(cell, bg, align='right')
            elif col_name == 'PREDIKSI':
                if val == 'Penerima':
                    cell.fill = PatternFill("solid", start_color="d1fae5")
                    cell.font = Font(name='Inter', size=10, bold=True, color='065f46')
                else:
                    cell.fill = PatternFill("solid", start_color="fecaca")
                    cell.font = Font(name='Inter', size=10, bold=True, color='991b1b')
                cell.alignment = Alignment(horizontal='center', vertical='center')
                thin = Side(style='thin', color='D9D9D9')
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            elif col_name == 'PENDAPATAN ORANG TUA':
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
                _style_data(cell, bg, align='right')
            else:
                _style_data(cell, bg)
    
    for i, header in enumerate(headers, 1):
        col_letter = get_column_letter(i)
        max_length = len(str(header))
        if not df_result.empty:
            col_values = df_result[header].astype(str)
            if len(col_values) > 0:
                max_val_len = col_values.str.len().max()
                max_length = max(max_length, max_val_len if max_val_len else 0)
        adjusted_width = min(max_length + 4, 35)
        ws.column_dimensions[col_letter].width = adjusted_width
    
    ws2 = wb.create_sheet("Ringkasan")
    ws2.merge_cells('A1:B1')
    ws2['A1'].value = "RINGKASAN PREDIKSI"
    ws2['A1'].font = Font(name='Inter', bold=True, size=13, color='FFFFFF')
    ws2['A1'].fill = PatternFill("solid", start_color="065f46")
    ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws2.row_dimensions[1].height = 26
    
    penerima = (df_result['PREDIKSI'] == 'Penerima').sum() if 'PREDIKSI' in df_result.columns else 0
    tidak_penerima = (df_result['PREDIKSI'] == 'Tidak Penerima').sum() if 'PREDIKSI' in df_result.columns else 0
    total = len(df_result)
    
    ringkasan_data = [
        ('Total Data', total),
        ('Jumlah Penerima', int(penerima)),
        ('Jumlah Tidak Penerima', int(tidak_penerima)),
        ('Persentase Penerima', f"{penerima/total*100:.1f}%" if total > 0 else "0%"),
        ('Persentase Tidak Penerima', f"{tidak_penerima/total*100:.1f}%" if total > 0 else "0%"),
    ]
    
    for i, (label, value) in enumerate(ringkasan_data, 3):
        cell_label = ws2.cell(row=i, column=1, value=label)
        _style_data(cell_label, align='left')
        cell_label.font = Font(name='Inter', size=10, bold=True)
        
        cell_value = ws2.cell(row=i, column=2, value=value)
        _style_data(cell_value)
    
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 22
    
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ─────────────────────────────────────────────────────────────
# FITUR BARU: PREDIKSI DATA BARU (DENGAN TAMBAHAN KE DATASET)
# ─────────────────────────────────────────────────────────────

def preprocess_single_input(kelas, pendapatan, pekerjaan, tanggungan, status_rumah, preprocess_info):
    """Preprocess single input untuk prediksi manual."""
    le_pekerjaan = preprocess_info['le_pekerjaan']
    scaler = preprocess_info['scaler']
    
    pekerjaan_enc = le_pekerjaan.transform([pekerjaan])[0] if pekerjaan in le_pekerjaan.classes_ else -1
    
    status_rumah_mapping = preprocess_info['status_rumah_mapping']
    status_rumah_enc = status_rumah_mapping.get(status_rumah, 0)
    
    pendapatan_zscore = (pendapatan - scaler.mean_[0]) / scaler.scale_[0]
    
    features = np.array([
        int(kelas),
        pendapatan_zscore,
        pekerjaan_enc,
        int(tanggungan),
        status_rumah_enc
    ]).reshape(1, -1)
    
    return features

def predict_single(model, features):
    """Melakukan prediksi untuk satu data."""
    prediction = model.predict(features)[0]
    probabilities = model.predict_proba(features)[0]
    return prediction, probabilities

def preprocess_batch_data(df_input, preprocess_info):
    """Preprocess data batch untuk prediksi massal."""
    required_columns = ['KELAS', 'PENDAPATAN ORANG TUA', 'PEKERJAAN ORANG TUA', 
                       'JUMLAH TANGGUNGAN', 'STATUS RUMAH']
    
    missing_cols = [col for col in required_columns if col not in df_input.columns]
    if missing_cols:
        raise ValueError(f"Kolom berikut tidak ditemukan: {', '.join(missing_cols)}")
    
    le_pekerjaan = preprocess_info['le_pekerjaan']
    scaler = preprocess_info['scaler']
    
    df_processed = df_input.copy()
    
    df_processed['KELAS_NUM'] = df_processed['KELAS'].astype(str).str.extract(r'(\d+)')
    df_processed['KELAS_NUM'] = pd.to_numeric(df_processed['KELAS_NUM'], errors='coerce').fillna(0).astype(int)
    
    pekerjaan_valid = set(le_pekerjaan.classes_)
    pekerjaan_input = set(df_processed['PEKERJAAN ORANG TUA'].astype(str).unique())
    pekerjaan_unknown = pekerjaan_input - pekerjaan_valid
    
    if pekerjaan_unknown:
        raise ValueError(
            f"Kategori pekerjaan tidak dikenal: {', '.join(pekerjaan_unknown)}. "
            f"Kategori yang valid: {', '.join(sorted(pekerjaan_valid))}"
        )
    
    df_processed['PEKERJAAN_ENC'] = le_pekerjaan.transform(
        df_processed['PEKERJAAN ORANG TUA'].astype(str)
    )
    
    status_rumah_mapping = preprocess_info['status_rumah_mapping']
    df_processed['STATUS_RUMAH_ENC'] = df_processed['STATUS RUMAH'].map(status_rumah_mapping)
    df_processed['STATUS_RUMAH_ENC'] = df_processed['STATUS_RUMAH_ENC'].fillna(0).astype(int)
    
    df_processed['PENDAPATAN_ZSCORE'] = (
        df_processed['PENDAPATAN ORANG TUA'] - scaler.mean_[0]
    ) / scaler.scale_[0]
    
    feature_cols = ['KELAS_NUM', 'PENDAPATAN_ZSCORE', 'PEKERJAAN_ENC', 
                   'JUMLAH TANGGUNGAN', 'STATUS_RUMAH_ENC']
    
    X = df_processed[feature_cols].values
    
    return df_input, X

def predict_batch(model, X, df_input):
    """Melakukan prediksi batch dan menambahkan hasil ke DataFrame."""
    predictions = model.predict(X)
    probabilities = model.predict_proba(X)
    
    df_result = df_input.copy()
    df_result['PREDIKSI'] = ['Penerima' if p == 1 else 'Tidak Penerima' for p in predictions]
    df_result['PROB_TIDAK (%)'] = (probabilities[:, 0] * 100).round(2)
    df_result['PROB_YA (%)'] = (probabilities[:, 1] * 100).round(2)
    
    return df_result

# ─────────────────────────────────────────────────────────────
# LANDING PAGE & MAIN APP
# ─────────────────────────────────────────────────────────────

def show_landing():
    """Halaman awal sebelum masuk dashboard."""
    
    st.markdown("""
    <div class="hero-home">
        <div style="position:absolute;top:18px;right:18px;width:60px;height:60px;
                    background:#FFE500;border:3px solid #000;transform:rotate(12deg);
                    box-shadow:3px 3px 0 #000;"></div>
        <div style="position:absolute;bottom:24px;right:80px;width:40px;height:40px;
                    border:3px solid #FFE500;border-radius:50%;"></div>
        <div style="position:absolute;top:50%;left:calc(100% - 160px);
                    font-family:'Bungee',cursive;font-size:6rem;color:rgba(255,255,255,0.06);
                    line-height:1;pointer-events:none;white-space:nowrap;">BSM</div>
        <div class="hero-badge">
            <span>🎓</span> MACHINE LEARNING &nbsp;·&nbsp; NAIVE BAYES &nbsp;·&nbsp; 2024
        </div>
        <div class="hero-title">Sistem Klasifikasi<br>Penerima BSM</div>
        <div class="hero-sub">
            Dashboard cerdas berbasis <strong>Gaussian Naive Bayes</strong> — membantu pengambilan 
            keputusan pemberian bantuan secara objektif, transparan, dan terukur.
        </div>
        <div style="display:flex; gap:12px; flex-wrap:wrap; margin-bottom:2rem;">
            <div style="background:#FFE500;border:3px solid #000;box-shadow:3px 3px 0 #000;
                        padding:12px 20px;color:#000;min-width:90px;text-align:center;">
                <div style="font-family:'Bungee',cursive;font-size:1.8rem;line-height:1;">5</div>
                <div style="font-family:'IBM Plex Mono',monospace;font-size:0.65rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:0.08em;margin-top:2px;">Fitur Input</div>
            </div>
            <div style="background:#00F5FF;border:3px solid #000;box-shadow:3px 3px 0 #000;
                        padding:12px 20px;color:#000;min-width:90px;text-align:center;">
                <div style="font-family:'Bungee',cursive;font-size:1.8rem;line-height:1;">2</div>
                <div style="font-family:'IBM Plex Mono',monospace;font-size:0.65rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:0.08em;margin-top:2px;">Kelas Output</div>
            </div>
            <div style="background:#B6FF00;border:3px solid #000;box-shadow:3px 3px 0 #000;
                        padding:12px 20px;color:#000;min-width:90px;text-align:center;">
                <div style="font-family:'Bungee',cursive;font-size:1.8rem;line-height:1;">GNB</div>
                <div style="font-family:'IBM Plex Mono',monospace;font-size:0.65rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:0.08em;margin-top:2px;">Algoritma</div>
            </div>
            <div style="background:#fff;border:3px solid #000;box-shadow:3px 3px 0 #000;
                        padding:12px 20px;color:#000;min-width:90px;text-align:center;">
                <div style="font-family:'Bungee',cursive;font-size:1.8rem;line-height:1;color:#2146FF;">100%</div>
                <div style="font-family:'IBM Plex Mono',monospace;font-size:0.65rem;font-weight:700;
                            text-transform:uppercase;letter-spacing:0.08em;margin-top:2px;">Open Source</div>
            </div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    col1, col2, col3 = st.columns(3)
    with col1:
        render_metric_card("🎯", "Objektif & Adil", "Probabilistik", "Klasifikasi berbasis probabilitas tanpa bias subjektif")
    with col2:
        render_metric_card("⚡", "Cepat & Ringan", "Real-time", "Training model dalam hitungan milidetik")
    with col3:
        render_metric_card("📊", "Transparan", "Terukur", "Semua metrik evaluasi divisualisasikan dengan jelas")

    st.markdown("<br>", unsafe_allow_html=True)

    st.markdown("### 📘 Penjelasan Metode & Rumus")
    with st.expander("🔍 Klik untuk melihat teori Gaussian Naive Bayes dan rumus lengkap", expanded=False):
        st.markdown("""
        **Gaussian Naive Bayes (GNB)** adalah algoritma klasifikasi probabilistik yang mengasumsikan 
        setiap fitur mengikuti distribusi normal (Gaussian).
        """)
        
        col_a, col_b = st.columns(2)
        with col_a:
            st.markdown("#### 🧠 Teorema Bayes")
            st.latex(r"P(C_k|X) = \frac{P(X|C_k) \cdot P(C_k)}{P(X)}")
            st.markdown("#### 📈 Gaussian PDF")
            st.latex(r"P(x_i|C_k) = \frac{1}{\sqrt{2\pi\sigma_k^2}} \exp\left(-\frac{(x_i-\mu_k)^2}{2\sigma_k^2}\right)")
            st.markdown("#### 📐 Z-Score")
            st.latex(r"z = \frac{x - \mu}{\sigma}")
        with col_b:
            st.markdown("#### 🎯 Accuracy")
            st.latex(r"\text{Accuracy} = \frac{TP + TN}{TP + TN + FP + FN}")
            st.markdown("#### 📊 Precision")
            st.latex(r"\text{Precision} = \frac{TP}{TP + FP}")
            st.markdown("#### 🔍 Recall")
            st.latex(r"\text{Recall} = \frac{TP}{TP + FN}")
            st.markdown("#### ⚖️ F1-Score")
            st.latex(r"F_1 = 2 \cdot \frac{\text{Precision} \cdot \text{Recall}}{\text{Precision} + \text{Recall}}")
    
    st.markdown("<br>", unsafe_allow_html=True)
    
    st.markdown("### 🔄 Alur Kerja Sistem")
    steps = [
        ("📂", "Upload Dataset", "Unggah file Excel (.xlsx) data BSM"),
        ("✏️", "Edit Data", "Ubah, tambah, atau hapus data langsung dari dashboard"),
        ("⚙️", "Preprocessing", "Missing value, duplikasi, encoding, standardisasi (No Data Leakage)"),
        ("📊", "Visualisasi", "Eksplorasi data dengan grafik Plotly"),
        ("🤖", "Training Model", "Latih Gaussian Naive Bayes dengan data train"),
        ("📈", "Evaluasi", "Accuracy, Precision, Recall, F1-Score, Classification Report"),
        ("🔮", "Prediksi Baru", "Prediksi manual / batch dari file Excel"),
        ("💾", "Export", "Unduh hasil ke file Excel profesional"),
    ]
    for icon, title, desc in steps:
        st.markdown(f"""
        <div style="display:flex; align-items:center; gap:14px; padding:12px 16px;
                    background:#fff; border:3px solid #000; box-shadow:4px 4px 0 #000;
                    margin-bottom:8px; transition:transform 0.15s;">
            <div style="font-size:1.6rem; width:44px; text-align:center;
                        background:#FFE500; border:2px solid #000; padding:6px;
                        box-shadow:2px 2px 0 #000; flex-shrink:0;">{icon}</div>
            <div>
                <div style="font-family:'Archivo Black',sans-serif; font-weight:900; color:#000;
                            font-size:0.88rem; text-transform:uppercase; letter-spacing:0.04em;">{title}</div>
                <div style="font-family:'IBM Plex Mono',monospace; font-size:0.75rem; color:#555;
                            margin-top:2px;">{desc}</div>
            </div>
        </div>
        """, unsafe_allow_html=True)
    
    st.markdown("<br>", unsafe_allow_html=True)
    
    _, center_col, _ = st.columns([1, 1.5, 1])
    with center_col:
        if st.button("🚀 Mulai Menggunakan Sistem", use_container_width=True, type="primary"):
            st.session_state.app_state = 'main'
            st.rerun()
    
    st.markdown("""
    <div class="footer">
        <div class="footer-brand">🎓 SISTEM KLASIFIKASI PENERIMA BSM</div>
        <div class="footer-divider"></div>
        <div style="font-family:'IBM Plex Mono',monospace;font-size:0.8rem;letter-spacing:0.06em;">
            GAUSSIAN NAIVE BAYES &nbsp;·&nbsp; MACHINE LEARNING &nbsp;·&nbsp; STREAMLIT
        </div>
        <div style="margin-top:0.5rem;font-family:'IBM Plex Mono',monospace;
                    color:rgba(255,255,255,0.4);font-size:0.7rem;letter-spacing:0.08em;">
            // DIBUAT UNTUK KEPERLUAN AKADEMIK DAN PENELITIAN //
        </div>
    </div>
    """, unsafe_allow_html=True)

def show_main():
    """Dashboard utama dengan sidebar navigasi."""
    
    with st.sidebar:
        st.markdown("""
        <div style="text-align:center; padding:1.4rem 0 1rem;">
            <div style="font-size:2.8rem; filter: drop-shadow(3px 3px 0 #000);">🎓</div>
            <div style="font-family:'Bungee',cursive; font-size:1.1rem; margin-top:6px; color:#FFE500;
                        letter-spacing:0.08em; text-shadow: 2px 2px 0 #000;">KLASIFIKASI BSM</div>
            <div style="font-family:'IBM Plex Mono',monospace; font-size:0.65rem; color:rgba(255,255,255,0.7);
                        margin-top:4px; letter-spacing:0.12em; text-transform:uppercase;">// NAIVE BAYES //</div>
            <div style="width:60%; height:3px; background:#FFE500; margin:10px auto 0;
                        border:1px solid #000;"></div>
        </div>
        """, unsafe_allow_html=True)
        st.markdown("---")
        
        page = st.radio(
            "Navigasi",
            [
                "🏠 Dashboard",
                "📂 Data & Preprocessing",
                "✏️ Edit Dataset",
                "📊 Visualisasi",
                "🤖 Training Model",
                "📈 Evaluasi",
                "🗂️ Hasil Prediksi",
                "🔮 Prediksi Data Baru",
                "📋 Kesimpulan Hasil",
                "💾 Export Hasil",
                "ℹ️ Tentang Sistem",
            ],
            label_visibility="collapsed"
        )
        
        st.markdown("---")
        st.markdown("**Status Sistem**")
        ds_loaded = st.session_state.df_raw is not None
        model_done = st.session_state.train_results is not None
        
        st.markdown(f"{'🟢' if ds_loaded else '⚪'} Dataset {'siap' if ds_loaded else 'belum dimuat'}")
        st.markdown(f"{'🟢' if model_done else '⚪'} Model {'terlatih' if model_done else 'belum dilatih'}")
        
        if ds_loaded:
            st.markdown("---")
            st.markdown("**Pengaturan Split Data**")
            
            if st.session_state.training_locked:
                test_size = st.slider(
                    "Ukuran Data Uji (%)", 0, 100, int(st.session_state.test_size * 100), 5,
                    disabled=True, help="Training sudah dikunci. Klik 'Reset Training' untuk mengubah."
                ) / 100.0
            else:
                test_size = st.slider(
                    "Ukuran Data Uji (%)", 0, 100, int(st.session_state.test_size * 100), 5,
                    help="Geser untuk mengatur proporsi data testing"
                ) / 100.0
                st.session_state.test_size = test_size
            
            if test_size == 0.0:
                st.info("ℹ️ Testing 0%: Semua data untuk training")
            elif test_size == 1.0:
                st.warning("⚠️ Testing 100%: Semua data untuk testing (tidak valid untuk training)")
            else:
                st.info(f"ℹ️ Split: {(1-test_size)*100:.0f}% Training | {test_size*100:.0f}% Testing")
        
        st.markdown("---")
        if st.button("🏠 Kembali ke Halaman Awal", use_container_width=True):
            st.session_state.app_state = 'home'
            st.rerun()

    # ═════════════════════════════════════════════════════════
    # PAGE ROUTING - DASHBOARD
    # ═════════════════════════════════════════════════════════

    if page == "🏠 Dashboard":
        render_sticky_header("🏠", "Dashboard Utama")
        
        if st.session_state.df_raw is not None:
            # Tampilkan change log ringkasan
            log = st.session_state.data_change_log
            col1, col2, col3, col4, col5 = st.columns(5)
            with col1:
                st.metric("Data Awal", log['initial_count'])
            with col2:
                st.metric("Data Saat Ini", log['current_count'])
            with col3:
                st.metric("Ditambahkan", log['added_count'], delta=f"+{log['added_count']}")
            with col4:
                st.metric("Dihapus", log['deleted_count'], delta=f"-{log['deleted_count']}")
            with col5:
                st.metric("Diedit", log['edited_count'])
            
            if log['last_activity']:
                st.caption(f"🕐 Aktivitas Terakhir: {log['last_activity']}")
            
            st.markdown("---")
            
            if st.session_state.df_processed is not None and st.session_state.preprocess_info:
                # Ambil data train dari info preprocessing
                df_train = st.session_state.preprocess_info.get('df_train', pd.DataFrame())
                df_test = st.session_state.preprocess_info.get('df_test', pd.DataFrame())
                
                n_train = len(df_train) if df_train is not None else 0
                n_test = len(df_test) if df_test is not None else 0
                n_total = n_train + n_test
                
                if n_train > 0 and 'LABEL_ENC' in df_train.columns:
                    n_ya_train = int((df_train['LABEL_ENC'] == 1).sum())
                    n_tdk_train = n_train - n_ya_train
                else:
                    n_ya_train = n_tdk_train = 0
                
                render_metric_cards([
                    {"icon": "🗃️", "label": "Total Data", "value": str(n_total), "desc": "setelah preprocessing"},
                    {"icon": "🏋️", "label": "Data Training", "value": str(n_train), "desc": f"{(n_train/n_total*100 if n_total>0 else 0):.1f}%"},
                    {"icon": "🧪", "label": "Data Testing", "value": str(n_test), "desc": f"{(n_test/n_total*100 if n_total>0 else 0):.1f}%"},
                    {"icon": "✅", "label": "Penerima (Train)", "value": str(n_ya_train), "desc": f"{n_ya_train/n_train*100 if n_train>0 else 0:.1f}% dari train"},
                ])
                
                # Warning untuk imbalanced dataset
                if n_train > 0 and n_ya_train > 0 and n_tdk_train > 0:
                    min_class_pct = min(n_ya_train, n_tdk_train) / n_train * 100
                    if min_class_pct < 20:
                        alert(f"⚠️ Dataset tidak seimbang (imbalanced)! Kelas minoritas hanya {min_class_pct:.1f}%. Akurasi dapat menjadi bias.", "warning")
            else:
                st.info("👋 Selamat datang! Silakan upload dataset di menu **📂 Data & Preprocessing**.")
        else:
            st.info("👋 Selamat datang! Silakan upload dataset di menu **📂 Data & Preprocessing**.")

    # ═════════════════════════════════════════════════════════
    # PAGE: DATA & PREPROCESSING
    # ═════════════════════════════════════════════════════════

    elif page == "📂 Data & Preprocessing":
        render_sticky_header("📂", "Upload & Preprocessing Dataset")
        
        col1, col2 = st.columns([3, 1])
        with col2:
            st.markdown("### 📥 Template")
            template_bytes = get_template_download()
            st.download_button(
                label="📄 Download Template Dataset",
                data=template_bytes,
                file_name="template_dataset_bsm.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                help="Download template Excel untuk mengisi data BSM"
            )
        
        uploaded = st.file_uploader(
            "Upload file Excel (.xlsx) data BSM",
            type=["xlsx"],
            help="Format: .xlsx | Header di baris ke-2 | Kolom sesuai template BSM"
        )
        
        if uploaded is not None:
            try:
                df_raw = pd.read_excel(uploaded, sheet_name=0, header=1)
                if len(df_raw.columns) < 14:
                    uploaded.seek(0)
                    df_raw = pd.read_excel(uploaded, sheet_name=0)
                df_raw = df_raw.dropna(how='all')
                st.session_state.df_raw = df_raw
                st.session_state.df_original = df_raw.copy()
                st.session_state.data_change_log['initial_count'] = len(df_raw)
                st.session_state.data_change_log['current_count'] = len(df_raw)
                
                st.success(f"✅ File berhasil dimuat: **{len(df_raw)}** baris × **{len(df_raw.columns)}** kolom")
                
                st.markdown("### 📋 Preview Dataset Lengkap")
                st.dataframe(df_raw, use_container_width=True, height=500)
                
                st.markdown("---")
                
                col_btn, col_info = st.columns([1, 3])
                with col_btn:
                    if st.button("⚙️ Jalankan Preprocessing", use_container_width=True, type="primary"):
                        with st.spinner("🔄 Memproses data (tanpa data leakage)..."):
                            progress_bar = st.progress(0)
                            for i in range(1, 101):
                                time.sleep(0.005)
                                progress_bar.progress(i)
                            df_p, info = run_preprocessing_correct(df_raw, st.session_state.test_size)
                            st.session_state.df_processed = df_p
                            st.session_state.preprocess_info = info
                            st.session_state.training_locked = False
                            st.session_state.train_results = None
                            st.session_state.batch_prediction_df = None
                        st.success("✅ Preprocessing selesai!")
                        st.rerun()
                
                if st.session_state.df_processed is not None and st.session_state.preprocess_info:
                    st.markdown("---")
                    st.markdown("### ✅ Hasil Preprocessing")
                    
                    info = st.session_state.preprocess_info
                    col_a, col_b, col_c, col_d = st.columns(4)
                    with col_a:
                        st.metric("Data Awal", info['total_rows'] + info['dup_removed'])
                    with col_b:
                        st.metric("Duplikat Dihapus", info['dup_removed'])
                    with col_c:
                        st.metric("Data Final", info['total_rows'])
                    with col_d:
                        st.metric("Split Data", f"Train: {info['train_rows']} | Test: {info['test_rows']}")
                    
                    st.dataframe(st.session_state.df_processed, use_container_width=True, height=500)
                    
            except Exception as e:
                st.error(f"❌ Error membaca file: {e}")
        
        with st.expander("📌 Panduan Format File", expanded=False):
            st.markdown("""
            | Kolom | Tipe Data | Deskripsi |
            |-------|-----------|-----------|
            | NO | integer | Nomor urut data |
            | NAMA PESERTA DIDIK | string | Nama lengkap siswa |
            | SEKOLAH | string | Nama sekolah |
            | KELAS | string | Tingkat kelas (7, 8, atau 9) |
            | NIK | string | Nomor Induk Kependudukan |
            | NISN | string | Nomor Induk Siswa Nasional |
            | NAMA AYAH/IBU | string | Nama orang tua |
            | KECAMATAN | string | Kecamatan tempat tinggal |
            | BESARAN BIAYA | integer | Besaran biaya (Rp) |
            | PENDAPATAN ORANG TUA | float | Pendapatan orang tua (Rp) |
            | PEKERJAAN ORANG TUA | string | Jenis pekerjaan orang tua |
            | JUMLAH TANGGUNGAN | integer | Jumlah tanggungan keluarga |
            | STATUS RUMAH | string | Milik Sendiri / Kontrak/sewa |
            | LABEL | string | Ya / Tidak |
            """)

    # ═════════════════════════════════════════════════════════
    # PAGE: EDIT DATASET (DIPERBAIKI - FIX DISABLED PARAMETER)
    # ═════════════════════════════════════════════════════════

    elif page == "✏️ Edit Dataset":
        render_sticky_header("✏️", "Edit Dataset")
        
        if st.session_state.df_raw is None:
            st.warning("⚠️ Dataset belum dimuat. Silakan upload dataset terlebih dahulu di menu **📂 Data & Preprocessing**.")
        else:
            st.markdown("### 📝 Edit Data Langsung")
            st.info("💡 Anda dapat mengedit data langsung pada tabel di bawah. Perubahan akan tersimpan di session.")
            
            df_editable = st.session_state.df_raw.copy()
            
            edited_df = st.data_editor(
                df_editable,
                use_container_width=True,
                height=500,
                num_rows="dynamic",
                key="data_editor"
            )
            
            col1, col2, col3 = st.columns(3)
            with col1:
                if st.button("💾 Simpan Perubahan", use_container_width=True, type="primary"):
                    if save_edited_data(edited_df):
                        st.success("✅ Perubahan berhasil disimpan!")
                        st.rerun()
                    else:
                        st.error("❌ Gagal menyimpan perubahan")
            
            with col2:
                if st.button("🔄 Reset ke Data Awal", use_container_width=True):
                    if reset_to_original():
                        st.success("✅ Data berhasil direset ke awal!")
                        st.rerun()
                    else:
                        st.error("❌ Gagal mereset data")
            
            with col3:
                st.markdown(f"**Jumlah Data:** {len(st.session_state.df_raw)} baris")
            
            st.markdown("---")
            st.markdown("### 🗑️ Hapus Data")
            
            # Tampilkan checkbox untuk memilih baris yang akan dihapus
            df_for_delete = st.session_state.df_raw.copy()
            if 'NO' in df_for_delete.columns:
                df_for_delete['PILIH'] = False
                display_cols = ['PILIH', 'NO', 'NAMA PESERTA DIDIK', 'KELAS', 'PENDAPATAN ORANG TUA', 'LABEL']
            else:
                df_for_delete['PILIH'] = False
                display_cols = ['PILIH'] + [c for c in df_for_delete.columns if c not in ['PILIH']][:5]
            
            # PERBAIKAN: disabled harus berupa list atau tuple, bukan None
            disabled_cols = []
            if 'NO' in df_for_delete.columns:
                disabled_cols = ['NO', 'NAMA PESERTA DIDIK', 'KELAS', 'PENDAPATAN ORANG TUA', 'LABEL']
            
            selected_df = st.data_editor(
                df_for_delete[display_cols],
                use_container_width=True,
                height=300,
                key="delete_selector",
                disabled=disabled_cols if disabled_cols else False  # False berarti tidak ada yang di-disable
            )
            
            selected_indices = selected_df[selected_df['PILIH'] == True].index.tolist()
            
            col_del1, col_del2 = st.columns([1, 3])
            with col_del1:
                if st.button("🗑️ Hapus Data Terpilih", use_container_width=True, type="secondary"):
                    if selected_indices:
                        success, deleted_count = delete_selected_rows(selected_indices)
                        if success:
                            st.success(f"✅ Berhasil menghapus {deleted_count} baris data!")
                            st.rerun()
                        else:
                            st.error("❌ Gagal menghapus data")
                    else:
                        st.warning("⚠️ Silakan pilih data yang akan dihapus terlebih dahulu")
            
            with col_del2:
                st.caption(f"Terpilih {len(selected_indices)} baris untuk dihapus")
            
            st.markdown("---")
            st.markdown("### 📋 Log Perubahan Data")
            
            log = st.session_state.data_change_log
            if log['activities']:
                log_df = pd.DataFrame(log['activities'])
                st.dataframe(log_df, use_container_width=True, hide_index=True)
            else:
                st.info("Belum ada aktivitas perubahan data")

    # ═════════════════════════════════════════════════════════
    # PAGE: VISUALISASI
    # ═════════════════════════════════════════════════════════

    elif page == "📊 Visualisasi":
        render_sticky_header("📊", "Visualisasi Data")
        
        if st.session_state.df_processed is None:
            st.warning("⚠️ Dataset belum diproses.")
        else:
            df = st.session_state.df_processed
            th = make_plotly_theme()
            colors = make_plotly_colors()
            
            tabs = st.tabs(["🥧 Distribusi Label", "💰 Pendapatan", "💼 Pekerjaan", "🏠 Status Rumah", "🎓 Kelas", "📊 Korelasi"])
            
            with tabs[0]:
                col1, col2 = st.columns([1, 1])
                with col1:
                    if '_DATASET_TYPE' in df.columns:
                        df_train_viz = df[df['_DATASET_TYPE'] == 'TRAIN'] if len(df[df['_DATASET_TYPE'] == 'TRAIN']) > 0 else df
                        label_counts_train = df_train_viz['LABEL'].value_counts()
                        fig = px.pie(
                            names=label_counts_train.index, 
                            values=label_counts_train.values,
                            color_discrete_sequence=[colors['penerima'], colors['tidak_penerima']], 
                            hole=0.45,
                            title="Distribusi Label Penerima BSM (Data TRAIN)"
                        )
                        fig.update_traces(
                            textposition='outside', 
                            textinfo='label+percent+value', 
                            textfont_size=13,
                            textfont_color='#FFFFFF'
                        )
                        fig.update_layout(**th, height=420)
                        st.plotly_chart(fig, use_container_width=True)
                    else:
                        label_counts = df['LABEL'].value_counts()
                        fig = px.pie(
                            names=label_counts.index, 
                            values=label_counts.values,
                            color_discrete_sequence=[colors['penerima'], colors['tidak_penerima']], 
                            hole=0.45,
                            title="Distribusi Label Penerima BSM"
                        )
                        fig.update_traces(
                            textposition='outside', 
                            textinfo='label+percent+value', 
                            textfont_size=13,
                            textfont_color='#FFFFFF'
                        )
                        fig.update_layout(**th, height=420)
                        st.plotly_chart(fig, use_container_width=True)
                with col2:
                    st.markdown("<br>", unsafe_allow_html=True)
                    col_met1, col_met2 = st.columns(2)
                    with col_met1:
                        if '_DATASET_TYPE' in df.columns:
                            st.metric("Total Penerima (Train)", int((df[df['_DATASET_TYPE']=='TRAIN']['LABEL_ENC']==1).sum()))
                        else:
                            st.metric("Total Penerima", int((df['LABEL_ENC']==1).sum()))
                    with col_met2:
                        if '_DATASET_TYPE' in df.columns:
                            st.metric("Total Tidak Penerima (Train)", int((df[df['_DATASET_TYPE']=='TRAIN']['LABEL_ENC']==0).sum()))
                        else:
                            st.metric("Total Tidak Penerima", int((df['LABEL_ENC']==0).sum()))
                    if '_DATASET_TYPE' in df.columns:
                        df_train_only = df[df['_DATASET_TYPE']=='TRAIN']
                        st.metric("Rasio (Train)", f"{(df_train_only['LABEL_ENC']==1).sum()/len(df_train_only)*100:.1f}% Penerima")
                    else:
                        st.metric("Rasio", f"{(df['LABEL_ENC']==1).sum()/len(df)*100:.1f}% Penerima")
            
            with tabs[1]:
                col1, col2 = st.columns(2)
                with col1:
                    fig1 = px.histogram(
                        df, x='PENDAPATAN ORANG TUA', nbins=30, 
                        color='_DATASET_TYPE' if '_DATASET_TYPE' in df.columns else None,
                        color_discrete_sequence=[colors['primary']],
                        title="Histogram Pendapatan (Rupiah)"
                    )
                    fig1.update_layout(**th, height=380)
                    st.plotly_chart(fig1, use_container_width=True)
                with col2:
                    fig2 = px.histogram(
                        df, x='PENDAPATAN_ZSCORE', nbins=30, 
                        color='_DATASET_TYPE' if '_DATASET_TYPE' in df.columns else None,
                        color_discrete_sequence=[colors['info']],
                        title="Histogram Pendapatan (Z-Score)"
                    )
                    fig2.update_layout(**th, height=380)
                    st.plotly_chart(fig2, use_container_width=True)
                
                fig3 = px.box(
                    df, x='LABEL', y='PENDAPATAN ORANG TUA', color='LABEL',
                    color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                    title="Box Plot Pendapatan per Status Penerima"
                )
                fig3.update_layout(**th, height=380)
                st.plotly_chart(fig3, use_container_width=True)
            
            with tabs[2]:
                pek_cnt = df['PEKERJAAN ORANG TUA'].value_counts().reset_index()
                pek_cnt.columns = ['Pekerjaan', 'Jumlah']
                fig = px.bar(
                    pek_cnt, y='Pekerjaan', x='Jumlah', orientation='h',
                    color='Jumlah', 
                    color_continuous_scale=colors['gradient_blue'],
                    title="Frekuensi Pekerjaan Orang Tua", 
                    text='Jumlah'
                )
                fig.update_traces(textposition='outside', textfont_color='#FFFFFF')
                fig.update_layout(**th, height=max(350, len(pek_cnt)*38),
                                 coloraxis_showscale=False, yaxis_categoryorder='total ascending')
                st.plotly_chart(fig, use_container_width=True)
                
                pek_label = df.groupby(['PEKERJAAN ORANG TUA', 'LABEL']).size().reset_index(name='Jumlah')
                fig2 = px.bar(
                    pek_label, y='PEKERJAAN ORANG TUA', x='Jumlah', color='LABEL',
                    orientation='h', 
                    color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                    barmode='stack', 
                    title="Pekerjaan vs Status Penerima"
                )
                fig2.update_layout(**th, height=max(350, len(pek_cnt)*38), yaxis_categoryorder='total ascending')
                st.plotly_chart(fig2, use_container_width=True)
            
            with tabs[3]:
                col1, col2 = st.columns(2)
                with col1:
                    sr_cnt = df['STATUS RUMAH'].value_counts()
                    fig1 = px.pie(
                        names=sr_cnt.index, values=sr_cnt.values,
                        color_discrete_sequence=[colors['primary'], colors['info']],
                        hole=0.4, 
                        title="Distribusi Status Rumah"
                    )
                    fig1.update_traces(textposition='inside', textinfo='label+percent', textfont_color='#FFFFFF')
                    fig1.update_layout(**th, height=350)
                    st.plotly_chart(fig1, use_container_width=True)
                with col2:
                    sr_label = df.groupby(['STATUS RUMAH', 'LABEL']).size().reset_index(name='Jumlah')
                    fig2 = px.bar(
                        sr_label, x='STATUS RUMAH', y='Jumlah', color='LABEL',
                        barmode='group', 
                        color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                        title="Status Rumah vs Status Penerima"
                    )
                    fig2.update_layout(**th, height=350)
                    st.plotly_chart(fig2, use_container_width=True)
            
            with tabs[4]:
                kl_cnt = df.groupby(['KELAS', 'LABEL']).size().reset_index(name='Jumlah')
                fig = px.bar(
                    kl_cnt, x='KELAS', y='Jumlah', color='LABEL', barmode='group',
                    color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                    title="Distribusi Kelas berdasarkan Status Penerima"
                )
                fig.update_layout(**th, height=400)
                st.plotly_chart(fig, use_container_width=True)
                
                fig2 = px.histogram(
                    df, x='JUMLAH TANGGUNGAN', color='LABEL',
                    barmode='overlay', opacity=0.75,
                    color_discrete_map={'Ya': colors['penerima'], 'Tidak': colors['tidak_penerima']},
                    title="Distribusi Jumlah Tanggungan per Status Penerima"
                )
                fig2.update_layout(**th, height=380)
                st.plotly_chart(fig2, use_container_width=True)
            
            with tabs[5]:
                num_cols = ['KELAS_NUM','PENDAPATAN_ZSCORE','PEKERJAAN ORANG TUA_ENC',
                           'JUMLAH TANGGUNGAN','STATUS RUMAH_ENC','LABEL_ENC']
                num_avail = [c for c in num_cols if c in df.columns]
                if len(num_avail) > 1:
                    corr = df[num_avail].corr().round(3)
                    fig = px.imshow(
                        corr, text_auto=True, 
                        color_continuous_scale=colors['gradient_blue'],
                        title="Heatmap Korelasi Fitur", 
                        aspect='auto'
                    )
                    fig.update_layout(**th, height=500)
                    st.plotly_chart(fig, use_container_width=True)

    # ═════════════════════════════════════════════════════════
    # PAGE: TRAINING MODEL
    # ═════════════════════════════════════════════════════════

    elif page == "🤖 Training Model":
        render_sticky_header("🤖", "Training Model")
        
        if st.session_state.df_processed is None:
            st.warning("⚠️ Dataset belum diproses.")
        elif st.session_state.preprocess_info is None:
            st.warning("⚠️ Preprocessing belum dilakukan.")
        else:
            info = st.session_state.preprocess_info
            df_train = info.get('df_train', pd.DataFrame())
            df_test = info.get('df_test', pd.DataFrame())
            
            n_train = len(df_train) if df_train is not None else 0
            n_test = len(df_test) if df_test is not None else 0
            n_valid = n_train + n_test
            
            if n_train > 0 and 'LABEL_ENC' in df_train.columns:
                n_ya_train = int((df_train['LABEL_ENC'] == 1).sum())
                n_tdk_train = n_train - n_ya_train
                min_class_pct = min(n_ya_train, n_tdk_train) / n_train * 100 if n_train > 0 else 0
            else:
                n_ya_train = n_tdk_train = 0
                min_class_pct = 0
            
            render_metric_cards([
                {"icon": "📦", "label": "Total Data Valid", "value": str(n_valid)},
                {"icon": "🏋️", "label": "Data Training", "value": str(n_train), "desc": f"{(n_train/n_valid*100 if n_valid>0 else 0):.0f}%"},
                {"icon": "🧪", "label": "Data Testing", "value": str(n_test), "desc": f"{(n_test/n_valid*100 if n_valid>0 else 0):.0f}%"},
                {"icon": "🔬", "label": "Algoritma", "value": "GNB", "desc": "Gaussian Naive Bayes"},
            ])
            
            if n_train > 0:
                st.markdown("---")
                st.markdown("### 📊 Distribusi Kelas pada Data Training")
                col1, col2, col3 = st.columns(3)
                with col1:
                    st.metric("Penerima BSM (Ya)", n_ya_train)
                with col2:
                    st.metric("Tidak Penerima (Tidak)", n_tdk_train)
                with col3:
                    st.metric("Persentase Penerima", f"{n_ya_train/n_train*100:.1f}%")
                
                if min_class_pct < 20:
                    alert(f"⚠️ Dataset tidak seimbang (imbalanced)! Kelas minoritas hanya {min_class_pct:.1f}%. Akurasi dapat menjadi bias.", "warning")
            
            if st.session_state.test_size == 1.0:
                st.error("❌ Testing 100% tidak valid untuk training otomatis.")
                st.info("ℹ️ Silakan ubah split data di sidebar.")
            elif n_train == 0:
                st.error("❌ Tidak ada data training. Silakan upload dataset atau ubah split data.")
            
            st.markdown("---")
            
            col_btn1, col_btn2 = st.columns(2)
            with col_btn1:
                btn_disabled = (st.session_state.test_size == 1.0 or st.session_state.training_locked or n_train == 0)
                if st.button("🚀 Mulai Training Model", disabled=btn_disabled, use_container_width=True, type="primary"):
                    with st.spinner("🔄 Melatih model..."):
                        progress_bar = st.progress(0)
                        for i in range(1, 101):
                            time.sleep(0.005)
                            progress_bar.progress(i)
                        res = run_training_correct(df_train, df_test)
                        st.session_state.train_results = res
                        st.session_state.training_locked = True
                        st.session_state.batch_prediction_df = None
                    
                    if n_test > 0 and res['test_acc']:
                        st.success(f"✅ Model berhasil dilatih! Akurasi Testing: {res['test_acc']*100:.2f}%")
                    else:
                        st.success("✅ Model berhasil dilatih! (Testing 0%)")
                    st.rerun()
            
            with col_btn2:
                if st.button("🔄 Reset Training", use_container_width=True):
                    st.session_state.train_results = None
                    st.session_state.training_locked = False
                    st.session_state.batch_prediction_df = None
                    st.rerun()
            
            if st.session_state.train_results:
                res = st.session_state.train_results
                st.markdown("---")
                st.markdown("### 📊 Ringkasan Hasil Training")
                
                c1, c2, c3, c4 = st.columns(4)
                with c1:
                    st.metric("Akurasi Training", f"{res['train_acc']*100:.2f}%" if res['train_acc'] is not None else "N/A")
                with c2:
                    if res['test_acc'] is not None:
                        st.metric("Akurasi Testing", f"{res['test_acc']*100:.2f}%")
                    else:
                        st.metric("Akurasi Testing", "N/A")
                with c3:
                    if res['y_pred'].size > 0:
                        benar = (res['y_test'].values == res['y_pred']).sum()
                        st.metric("Prediksi Benar", int(benar))
                    else:
                        st.metric("Prediksi Benar", "N/A")
                with c4:
                    if res['y_pred'].size > 0:
                        salah = (res['y_test'].values != res['y_pred']).sum()
                        st.metric("Prediksi Salah", int(salah))
                    else:
                        st.metric("Prediksi Salah", "N/A")
                
                if res['model'] is not None:
                    with st.expander("🧮 Parameter Model (Theta & Variance)", expanded=False):
                        model = res['model']
                        fc = res['feature_cols']
                        st.dataframe(pd.DataFrame({
                            'Fitur': fc,
                            'Mean Tidak (0)': model.theta_[0].round(4),
                            'Mean Ya (1)': model.theta_[1].round(4),
                            'Std Tidak (0)': np.sqrt(model.var_[0]).round(4),
                            'Std Ya (1)': np.sqrt(model.var_[1]).round(4),
                        }), use_container_width=True, hide_index=True)

    # ═════════════════════════════════════════════════════════
    # PAGE: EVALUASI
    # ═════════════════════════════════════════════════════════

    elif page == "📈 Evaluasi":
        render_sticky_header("📈", "Evaluasi Model")
        
        if st.session_state.train_results is None:
            st.warning("⚠️ Model belum dilatih.")
        else:
            res = st.session_state.train_results
            th = make_plotly_theme()
            colors = make_plotly_colors()
            
            # Tampilkan classification report
            if res['classification_report']:
                st.markdown("### 📋 Classification Report (Scikit-learn)")
                report_df = pd.DataFrame(res['classification_report']).transpose()
                st.dataframe(report_df.round(4), use_container_width=True)
            
            st.markdown("---")
            
            render_metric_cards([
                {"icon": "🏋️", "label": "Akurasi Training", "value": f"{res['train_acc']*100:.2f}%" if res['train_acc'] is not None else "N/A"},
                {"icon": "🧪", "label": "Akurasi Testing", "value": f"{res['test_acc']*100:.2f}%" if res['test_acc'] is not None else "N/A"},
                {"icon": "📊", "label": "F1-Score Macro", "value": f"{res['macro_f1']*100:.2f}%" if res['macro_f1'] else "N/A"},
                {"icon": "⚖️", "label": "F1-Score Weighted", "value": f"{res['w_f1']*100:.2f}%" if res['w_f1'] else "N/A"},
            ])
            
            st.markdown("---")
            tabs = st.tabs(["📋 Tabel Metrik", "🗂️ Confusion Matrix", "📊 Grafik Evaluasi"])
            
            with tabs[0]:
                col1, col2 = st.columns(2)
                with col1:
                    st.markdown("**Per Kelas**")
                    st.dataframe(pd.DataFrame({
                        'Metrik': ['Precision', 'Recall', 'F1-Score', 'Support'],
                        'Tidak Penerima (0)': [
                            f"{res['prec0']:.4f}" if res['prec0'] else "N/A",
                            f"{res['rec0']:.4f}" if res['rec0'] else "N/A",
                            f"{res['f1_0']:.4f}" if res['f1_0'] else "N/A",
                            int(res['sup0']) if res['sup0'] else 0
                        ],
                        'Penerima (1)': [
                            f"{res['prec1']:.4f}" if res['prec1'] else "N/A",
                            f"{res['rec1']:.4f}" if res['rec1'] else "N/A",
                            f"{res['f1_1']:.4f}" if res['f1_1'] else "N/A",
                            int(res['sup1']) if res['sup1'] else 0
                        ],
                    }), use_container_width=True, hide_index=True)
                with col2:
                    st.markdown("**Average**")
                    st.dataframe(pd.DataFrame({
                        'Tipe': ['Macro', 'Weighted'],
                        'Precision': [f"{res['macro_prec']:.4f}" if res['macro_prec'] else "N/A", f"{res['w_prec']:.4f}" if res['w_prec'] else "N/A"],
                        'Recall': [f"{res['macro_rec']:.4f}" if res['macro_rec'] else "N/A", f"{res['w_rec']:.4f}" if res['w_rec'] else "N/A"],
                        'F1-Score': [f"{res['macro_f1']:.4f}" if res['macro_f1'] else "N/A", f"{res['w_f1']:.4f}" if res['w_f1'] else "N/A"],
                    }), use_container_width=True, hide_index=True)
            
            with tabs[1]:
                if res['cm'] is not None and res['cm'].sum() > 0:
                    cm = res['cm']
                    fig = go.Figure(data=go.Heatmap(
                        z=cm,
                        x=['Predicted: Tidak', 'Predicted: Ya'],
                        y=['Actual: Tidak', 'Actual: Ya'],
                        text=[[str(v) for v in row] for row in cm],
                        texttemplate="%{text}",
                        textfont={"size": 22, "family": "JetBrains Mono", "color": "#FFFFFF"},
                        colorscale=[[0, "#1E293B"], [0.5, "#3B82F6"], [1, "#93C5FD"]],
                        showscale=True, 
                        colorbar=dict(title="Jumlah", tickfont=dict(color="#FFFFFF"), title_font=dict(color="#FFFFFF"))
                    ))
                    fig.update_layout(**th, title="Confusion Matrix", height=420)
                    st.plotly_chart(fig, use_container_width=True)
                    
                    if cm.shape == (2, 2):
                        TP = cm[1, 1]
                        TN = cm[0, 0]
                        FP = cm[0, 1]
                        FN = cm[1, 0]
                        
                        col_a, col_b, col_c, col_d = st.columns(4)
                        with col_a:
                            st.metric("True Positive", TP)
                        with col_b:
                            st.metric("True Negative", TN)
                        with col_c:
                            st.metric("False Positive", FP)
                        with col_d:
                            st.metric("False Negative", FN)
            
            with tabs[2]:
                if res['train_acc'] is not None and res['test_acc'] is not None:
                    metrics = {
                        'Accuracy Train': res['train_acc'], 
                        'Accuracy Test': res['test_acc'],
                        'Precision (0)': res['prec0'] if res['prec0'] else 0, 
                        'Recall (0)': res['rec0'] if res['rec0'] else 0, 
                        'F1 (0)': res['f1_0'] if res['f1_0'] else 0,
                        'Precision (1)': res['prec1'] if res['prec1'] else 0, 
                        'Recall (1)': res['rec1'] if res['rec1'] else 0, 
                        'F1 (1)': res['f1_1'] if res['f1_1'] else 0,
                    }
                    fig1 = px.bar(
                        x=list(metrics.keys()), y=[v*100 for v in metrics.values()],
                        text=[f"{v*100:.1f}%" for v in metrics.values()],
                        color_discrete_sequence=[colors['primary']]*len(metrics),
                        title="Perbandingan Semua Metrik (%)"
                    )
                    fig1.update_traces(textfont_color='#FFFFFF')
                    fig1.update_layout(**th, yaxis_title="Nilai (%)", height=400)
                    st.plotly_chart(fig1, use_container_width=True)

    # ═════════════════════════════════════════════════════════
    # PAGE: HASIL PREDIKSI
    # ═════════════════════════════════════════════════════════

    elif page == "🗂️ Hasil Prediksi":
        render_sticky_header("🗂️", "Hasil Prediksi")
        
        if st.session_state.train_results is None:
            st.warning("⚠️ Model belum dilatih.")
        elif st.session_state.train_results['y_pred'].size == 0:
            st.warning("⚠️ Tidak ada hasil prediksi. Pastikan data testing tersedia.")
        else:
            res = st.session_state.train_results
            df_r = res['results_df']
            
            if len(df_r) > 0:
                benar = (df_r['Status'] == '✓ Benar').sum()
                salah = len(df_r) - benar
                
                render_metric_cards([
                    {"icon": "📋", "label": "Total Data Uji", "value": str(len(df_r))},
                    {"icon": "✓", "label": "Prediksi Benar", "value": str(benar), "desc": f"{benar/len(df_r)*100:.1f}%"},
                    {"icon": "✗", "label": "Prediksi Salah", "value": str(salah), "desc": f"{salah/len(df_r)*100:.1f}%"},
                    {"icon": "🎯", "label": "Akurasi", "value": f"{benar/len(df_r)*100:.2f}%"},
                ])
                
                st.markdown("---")
                col1, col2, col3 = st.columns(3)
                with col1:
                    filter_status = st.selectbox("Filter Status", ["Semua", "✓ Benar", "✗ Salah"])
                with col2:
                    filter_actual = st.selectbox("Filter Actual", ["Semua", "Penerima", "Tidak Penerima"])
                with col3:
                    filter_pred = st.selectbox("Filter Predicted", ["Semua", "Penerima", "Tidak Penerima"])
                
                df_disp = df_r.copy()
                if filter_status != "Semua":
                    df_disp = df_disp[df_disp['Status'] == filter_status]
                if filter_actual != "Semua":
                    df_disp = df_disp[df_disp['Actual'] == filter_actual]
                if filter_pred != "Semua":
                    df_disp = df_disp[df_disp['Predicted'] == filter_pred]
                
                st.caption(f"Menampilkan {len(df_disp)} dari {len(df_r)} prediksi")
                st.dataframe(df_disp, use_container_width=True, height=500, hide_index=True)
            else:
                st.info("Belum ada hasil prediksi.")

    # ═════════════════════════════════════════════════════════
    # PAGE: PREDIKSI DATA BARU
    # ═════════════════════════════════════════════════════════

    elif page == "🔮 Prediksi Data Baru":
        render_sticky_header("🔮", "Prediksi Data Baru")
        
        if st.session_state.train_results is None:
            st.warning("⚠️ Model belum dilatih. Silakan lakukan training model terlebih dahulu di menu **🤖 Training Model**.")
        elif st.session_state.preprocess_info is None:
            st.warning("⚠️ Data preprocessing belum tersedia. Silakan upload dan preprocessing data terlebih dahulu di menu **📂 Data & Preprocessing**.")
        else:
            model = st.session_state.train_results['model']
            preprocess_info = st.session_state.preprocess_info['preprocess_info']
            
            if model is None:
                st.warning("⚠️ Model tidak tersedia. Silakan lakukan training terlebih dahulu.")
            else:
                tabs = st.tabs(["✍️ Input Manual", "📂 Prediksi Massal"])
                
                with tabs[0]:
                    st.markdown("### ✍️ Prediksi Manual & Tambah ke Dataset")
                    st.markdown("Masukkan data siswa secara manual untuk memprediksi status penerimaan BSM.")
                    
                    with st.form("prediction_manual_form"):
                        col1, col2 = st.columns(2)
                        
                        with col1:
                            kelas = st.selectbox("📚 Kelas", options=["7", "8", "9"])
                            pendapatan = st.number_input("💰 Pendapatan Orang Tua (Rp)", min_value=0, value=2000000, step=100000, format="%d")
                            pekerjaan_options = list(preprocess_info['le_pekerjaan'].classes_)
                            pekerjaan = st.selectbox("💼 Pekerjaan Orang Tua", options=pekerjaan_options)
                        
                        with col2:
                            tanggungan = st.number_input("👨‍👩‍👧‍👦 Jumlah Tanggungan", min_value=0, max_value=20, value=3, step=1)
                            status_rumah = st.selectbox("🏠 Status Rumah", options=["Milik Sendiri", "Kontrak/sewa"])
                            label_aktual = st.selectbox("🏷️ Label Aktual (Opsional)", options=["", "Ya", "Tidak"], help="Isi jika ingin menambahkan ke dataset")
                        
                        col_btn1, col_btn2 = st.columns(2)
                        with col_btn1:
                            submitted = st.form_submit_button("🔮 Prediksi Sekarang", use_container_width=True, type="primary")
                        with col_btn2:
                            add_to_dataset = st.form_submit_button("➕ Prediksi & Tambahkan ke Dataset", use_container_width=True)
                    
                    if submitted or add_to_dataset:
                        with st.spinner("🔮 Melakukan prediksi..."):
                            time.sleep(0.5)
                            
                            features = preprocess_single_input(
                                kelas=kelas, pendapatan=pendapatan, pekerjaan=pekerjaan,
                                tanggungan=tanggungan, status_rumah=status_rumah, preprocess_info=preprocess_info
                            )
                            
                            prediction, probabilities = predict_single(model, features)
                            
                            st.markdown("---")
                            st.markdown("### 🎯 Hasil Prediksi")
                            
                            if prediction == 1:
                                card_class = "success"
                                icon = "✅"
                                result_text = "PENERIMA BSM"
                                badge_class = "badge-penerima"
                            else:
                                card_class = "danger"
                                icon = "❌"
                                result_text = "TIDAK PENERIMA BSM"
                                badge_class = "badge-tidak"
                            
                            st.markdown(f"""
                            <div class="prediction-card {card_class}">
                                <div class="prediction-icon">{icon}</div>
                                <div class="prediction-result">{result_text}</div>
                                <div style="margin: 1rem 0;">
                                    <span class="{badge_class}">{result_text}</span>
                                </div>
                            </div>
                            """, unsafe_allow_html=True)
                            
                            col_prob1, col_prob2 = st.columns(2)
                            with col_prob1:
                                st.metric("Probabilitas TIDAK", f"{probabilities[0]*100:.2f}%")
                            with col_prob2:
                                st.metric("Probabilitas YA", f"{probabilities[1]*100:.2f}%")
                            
                            # Jika user memilih tambah ke dataset
                            if add_to_dataset and label_aktual:
                                new_row = {
                                    'NO': None,
                                    'NAMA PESERTA DIDIK': '-',
                                    'SEKOLAH': '-',
                                    'KELAS': kelas,
                                    'NIK': '-',
                                    'NISN': '-',
                                    'NAMA AYAH/IBU': '-',
                                    'KECAMATAN': '-',
                                    'BESARAN BIAYA': 0,
                                    'PENDAPATAN ORANG TUA': pendapatan,
                                    'PEKERJAAN ORANG TUA': pekerjaan,
                                    'JUMLAH TANGGUNGAN': tanggungan,
                                    'STATUS RUMAH': status_rumah,
                                    'LABEL': label_aktual
                                }
                                
                                success, new_no = add_new_data_to_dataset(new_row)
                                if success:
                                    st.success(f"✅ Data berhasil ditambahkan ke dataset! (NO: {new_no})")
                                    st.info("🔄 Silakan lakukan preprocessing ulang dan training ulang untuk menggunakan data baru.")
                                else:
                                    st.error("❌ Gagal menambahkan data ke dataset")
                
                with tabs[1]:
                    st.markdown("### 📂 Prediksi Massal")
                    st.markdown("Upload file Excel berisi data siswa untuk diprediksi secara massal sekaligus.")
                    
                    with st.expander("📌 Panduan Format File untuk Prediksi Massal", expanded=False):
                        st.markdown("""
                        **Format file Excel (.xlsx) yang wajib dipenuhi:**
                        
                        | Kolom | Tipe Data | Keterangan | Contoh |
                        |-------|-----------|------------|--------|
                        | KELAS | string/angka | Tingkat kelas siswa | 7, 8, atau 9 |
                        | PENDAPATAN ORANG TUA | angka | Pendapatan bulanan (Rp) | 2000000 |
                        | PEKERJAAN ORANG TUA | string | Jenis pekerjaan | Buruh, Petani, dll |
                        | JUMLAH TANGGUNGAN | angka | Jumlah tanggungan keluarga | 3 |
                        | STATUS RUMAH | string | Status kepemilikan | "Milik Sendiri" / "Kontrak/sewa" |
                        """)
                    
                    uploaded_file = st.file_uploader(
                        "Upload File Excel untuk Prediksi Massal",
                        type=['xlsx'],
                        help="Upload file .xlsx dengan kolom sesuai panduan di atas"
                    )
                    
                    if uploaded_file is not None:
                        try:
                            df_input = pd.read_excel(uploaded_file)
                            st.success(f"✅ File berhasil dimuat: **{len(df_input)}** baris × **{len(df_input.columns)}** kolom")
                            
                            st.markdown("#### 📋 Preview Data Input (10 baris pertama)")
                            st.dataframe(df_input.head(10), use_container_width=True)
                            
                            if st.button("📊 Jalankan Prediksi Massal", use_container_width=True, type="primary", key="btn_batch_predict"):
                                with st.spinner("🔄 Memproses data dan melakukan prediksi..."):
                                    progress_bar = st.progress(0)
                                    for i in range(1, 101, 20):
                                        time.sleep(0.02)
                                        progress_bar.progress(i)
                                    
                                    try:
                                        df_input_copy, X = preprocess_batch_data(df_input, preprocess_info)
                                        df_result = predict_batch(model, X, df_input_copy)
                                        st.session_state.batch_prediction_df = df_result
                                        progress_bar.progress(100)
                                        st.success("✅ Prediksi massal selesai!")
                                        time.sleep(0.5)
                                        st.rerun()
                                    except ValueError as ve:
                                        st.error(f"❌ Error Validasi: {ve}")
                                        progress_bar.progress(100)
                                    except Exception as e:
                                        st.error(f"❌ Error: {e}")
                            
                            if st.session_state.batch_prediction_df is not None:
                                df_result = st.session_state.batch_prediction_df
                                
                                st.markdown("---")
                                st.markdown("### 📊 Ringkasan Hasil Prediksi")
                                
                                penerima = int((df_result['PREDIKSI'] == 'Penerima').sum())
                                tidak_penerima = int((df_result['PREDIKSI'] == 'Tidak Penerima').sum())
                                total = len(df_result)
                                
                                col_a, col_b, col_c = st.columns(3)
                                with col_a:
                                    st.metric("📦 Total Data", total)
                                with col_b:
                                    pct_penerima = penerima / total * 100 if total > 0 else 0
                                    st.metric("✅ Penerima BSM", penerima, delta=f"{pct_penerima:.1f}%")
                                with col_c:
                                    pct_tidak = tidak_penerima / total * 100 if total > 0 else 0
                                    st.metric("❌ Tidak Penerima", tidak_penerima, delta=f"{pct_tidak:.1f}%", delta_color="inverse")
                                
                                st.markdown("#### 📋 Tabel Hasil Prediksi")
                                st.dataframe(
                                    df_result, use_container_width=True, height=450, hide_index=True,
                                    column_config={
                                        'PROB_TIDAK (%)': st.column_config.NumberColumn(format="%.2f%%"),
                                        'PROB_YA (%)': st.column_config.NumberColumn(format="%.2f%%"),
                                    }
                                )
                                
                                st.markdown("---")
                                excel_bytes = export_batch_prediction_excel(df_result)
                                
                                st.download_button(
                                    label="📥 Download Hasil Prediksi (.xlsx)",
                                    data=excel_bytes,
                                    file_name="hasil_prediksi_bsm.xlsx",
                                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                    use_container_width=True
                                )
                                
                        except Exception as e:
                            st.error(f"❌ Error membaca file: {e}")

    # ═════════════════════════════════════════════════════════
    # PAGE: KESIMPULAN HASIL
    # ═════════════════════════════════════════════════════════

    elif page == "📋 Kesimpulan Hasil":
        render_sticky_header("📋", "Kesimpulan Hasil")
        
        if st.session_state.train_results is None:
            st.warning("⚠️ Model belum dilatih. Silakan lakukan training terlebih dahulu.")
        else:
            res = st.session_state.train_results
            
            st.markdown("## 🎯 Kesimpulan dan Interpretasi Model")
            st.markdown("---")
            
            col1, col2 = st.columns(2)
            with col1:
                render_metric_card("🤖", "Status Model", "Telah Dilatih", f"Algoritma: Gaussian Naive Bayes")
                render_metric_card("📊", "Jumlah Data Training", str(len(res['X_train'])), "Data yang digunakan untuk melatih model")
                if len(res['y_pred']) > 0:
                    render_metric_card("✅", "Jumlah Prediksi Benar", str(int((res['y_test'].values == res['y_pred']).sum())), "Dari data testing")
                else:
                    render_metric_card("✅", "Jumlah Prediksi Benar", "N/A", "Dari data testing")
            with col2:
                render_metric_card("🎯", "Akurasi Akhir", f"{res['test_acc']*100:.2f}%" if res['test_acc'] else "N/A", "Performa model pada data testing")
                render_metric_card("📊", "Jumlah Data Testing", str(len(res['X_test'])) if len(res['X_test']) > 0 else "0", "Data uji model")
                if len(res['y_pred']) > 0:
                    render_metric_card("❌", "Jumlah Prediksi Salah", str(int((res['y_test'].values != res['y_pred']).sum())), "Dari data testing")
                else:
                    render_metric_card("❌", "Jumlah Prediksi Salah", "N/A", "Dari data testing")
            
            st.markdown("---")
            
            st.markdown("## 📈 Ringkasan Performa Model")
            
            col_a, col_b, col_c = st.columns(3)
            with col_a:
                st.markdown("**Metrik Utama**")
                st.markdown(f"- **Akurasi Training:** {res['train_acc']*100:.2f}%" if res['train_acc'] else "- **Akurasi Training:** N/A")
                st.markdown(f"- **Akurasi Testing:** {res['test_acc']*100:.2f}%" if res['test_acc'] else "- **Akurasi Testing:** N/A")
                st.markdown(f"- **F1-Score Macro:** {res['macro_f1']*100:.2f}%" if res['macro_f1'] else "- **F1-Score Macro:** N/A")
                st.markdown(f"- **F1-Score Weighted:** {res['w_f1']*100:.2f}%" if res['w_f1'] else "- **F1-Score Weighted:** N/A")
            with col_b:
                st.markdown("**Kelas Tidak Penerima (0)**")
                st.markdown(f"- **Precision:** {res['prec0']*100:.2f}%" if res['prec0'] else "- **Precision:** N/A")
                st.markdown(f"- **Recall:** {res['rec0']*100:.2f}%" if res['rec0'] else "- **Recall:** N/A")
                st.markdown(f"- **F1-Score:** {res['f1_0']*100:.2f}%" if res['f1_0'] else "- **F1-Score:** N/A")
                st.markdown(f"- **Support:** {int(res['sup0'])}" if res['sup0'] else "- **Support:** 0")
            with col_c:
                st.markdown("**Kelas Penerima (1)**")
                st.markdown(f"- **Precision:** {res['prec1']*100:.2f}%" if res['prec1'] else "- **Precision:** N/A")
                st.markdown(f"- **Recall:** {res['rec1']*100:.2f}%" if res['rec1'] else "- **Recall:** N/A")
                st.markdown(f"- **F1-Score:** {res['f1_1']*100:.2f}%" if res['f1_1'] else "- **F1-Score:** N/A")
                st.markdown(f"- **Support:** {int(res['sup1'])}" if res['sup1'] else "- **Support:** 0")
            
            st.markdown("---")
            
            st.markdown("## 🔍 Interpretasi Hasil")
            
            if res['test_acc']:
                if res['test_acc'] >= 0.85:
                    st.success("✅ **Model memiliki performa yang sangat baik** dengan akurasi di atas 85%.")
                elif res['test_acc'] >= 0.70:
                    st.info("📊 **Model memiliki performa yang cukup baik** dengan akurasi antara 70-85%.")
                else:
                    st.warning("⚠️ **Model memiliki performa yang perlu ditingkatkan** dengan akurasi di bawah 70%.")
            
            if res['prec1'] and res['rec1']:
                st.markdown(f"""
                - **Kelas Penerima (Ya)**: Precision = {res['prec1']*100:.1f}%, Recall = {res['rec1']*100:.1f}%
                - **Kelas Tidak Penerima (Tidak)**: Precision = {res['prec0']*100:.1f}%, Recall = {res['rec0']*100:.1f}%
                """)
                
                if res['prec1'] > res['rec1']:
                    st.markdown("- Model cenderung **konservatif** dalam memprediksi kelas Penerima (lebih sedikit false positive)")
                elif res['rec1'] > res['prec1']:
                    st.markdown("- Model cenderung **agresif** dalam memprediksi kelas Penerima (lebih sedikit false negative)")
            
            st.markdown("---")
            
            st.markdown("## ⚠️ Keterbatasan Model")
            st.markdown("""
            1. **Asumsi Independensi Fitur**: Naive Bayes mengasumsikan semua fitur independen, padahal dalam praktiknya fitur-fitur seperti pendapatan dan pekerjaan orang tua mungkin saling berkorelasi.
            2. **Distribusi Normal**: Gaussian Naive Bayes mengasumsikan data numerik berdistribusi normal. Jika data tidak terdistribusi normal, performa model dapat menurun.
            3. **Imbalanced Dataset**: Jika dataset tidak seimbang, model dapat bias terhadap kelas mayoritas.
            4. **Keterbatasan Fitur**: Model hanya menggunakan 5 fitur. Penambahan fitur yang relevan dapat meningkatkan akurasi.
            """)
            
            st.markdown("---")
            
            st.markdown("## 💡 Rekomendasi Penggunaan Sistem")
            st.markdown("""
            1. **Gunakan sebagai alat bantu keputusan**, bukan sebagai penentu final. Keputusan final tetap harus mempertimbangkan aspek-aspek lain.
            2. **Perbaharui model secara berkala** dengan data baru untuk menjaga akurasi.
            3. **Pastikan data input berkualitas** dan sesuai dengan format yang ditentukan.
            4. **Lakukan validasi lapangan** untuk data yang diprediksi sebagai penerima maupun tidak penerima.
            5. **Gunakan prediksi batch** untuk efisiensi jika memiliki banyak data.
            """)
            
            st.markdown("---")
            
            # Tombol export ringkasan model
            if st.button("📥 Export Ringkasan Model ke Excel", use_container_width=True):
                excel_bytes = export_model_summary_excel(res, st.session_state.preprocess_info)
                st.download_button(
                    label="⬇️ Download Ringkasan Model (.xlsx)",
                    data=excel_bytes,
                    file_name="ringkasan_model_bsm.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

    # ═════════════════════════════════════════════════════════
    # PAGE: EXPORT HASIL
    # ═════════════════════════════════════════════════════════

    elif page == "💾 Export Hasil":
        render_sticky_header("💾", "Export & Download")
        
        if st.session_state.df_processed is None:
            st.warning("⚠️ Dataset belum diproses.")
        elif st.session_state.preprocess_info is None:
            st.warning("⚠️ Preprocessing info tidak tersedia.")
        else:
            df = st.session_state.df_processed
            info = st.session_state.preprocess_info
            le = info['preprocess_info']['le_pekerjaan']
            
            col1, col2 = st.columns(2)
            with col1:
                st.markdown("### 📊 Data Preprocessing")
                preproc_bytes = export_preprocessing_excel(df, le)
                st.download_button("⬇️ Download Data Preprocessing (.xlsx)", data=preproc_bytes,
                                  file_name="data_bsm_preprocessed.xlsx",
                                  mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                  use_container_width=True)
                
                st.markdown("---")
                st.markdown("### 📐 Data Normalisasi")
                norm_bytes = export_normalisasi_excel(df, info)
                st.download_button("⬇️ Download Hasil Normalisasi (.xlsx)", data=norm_bytes,
                                  file_name="data_normalisasi_zscore.xlsx",
                                  mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                  use_container_width=True)
            
            with col2:
                if st.session_state.train_results:
                    res = st.session_state.train_results
                    st.markdown("### 🎯 Hasil Prediksi")
                    if len(res['results_df']) > 0:
                        output = io.BytesIO()
                        with pd.ExcelWriter(output, engine='openpyxl') as writer:
                            res['results_df'].to_excel(writer, index=False, sheet_name='Hasil Prediksi')
                        output.seek(0)
                        st.download_button("⬇️ Download Hasil Prediksi (.xlsx)", data=output.getvalue(),
                                          file_name="hasil_prediksi_naive_bayes.xlsx",
                                          mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                          use_container_width=True)
                    else:
                        st.info("ℹ️ Tidak ada hasil prediksi untuk diekspor.")
                    
                    st.markdown("---")
                    st.markdown("### 📋 Ringkasan Model")
                    summary_bytes = export_model_summary_excel(res, info)
                    st.download_button("⬇️ Download Ringkasan Model (.xlsx)", data=summary_bytes,
                                      file_name="ringkasan_model_bsm.xlsx",
                                      mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                      use_container_width=True)
                else:
                    st.info("ℹ️ Latih model terlebih dahulu untuk mengekspor hasil prediksi dan ringkasan model.")
            
            st.markdown("---")
            st.markdown("### 📋 Template Dataset")
            template_bytes = get_template_download()
            st.download_button(
                label="📄 Download Template Dataset (.xlsx)",
                data=template_bytes,
                file_name="template_dataset_bsm.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                help="Download template Excel untuk mengisi data BSM"
            )

    # ═════════════════════════════════════════════════════════
    # PAGE: TENTANG SISTEM
    # ═════════════════════════════════════════════════════════

    elif page == "ℹ️ Tentang Sistem":
        render_sticky_header("ℹ️", "Tentang Sistem")
        
        col1, col2, col3 = st.columns(3)
        with col1:
            render_metric_card("🎓", "Tujuan", "Klasifikasi BSM", "Membantu keputusan pemberian BSM secara objektif")
        with col2:
            render_metric_card("🔬", "Algoritma", "Gaussian NB", "Probabilistik berbasis Teorema Bayes")
        with col3:
            render_metric_card("📊", "Fitur", "5 Input", "Kelas, Pendapatan, Pekerjaan, Tanggungan, Status Rumah")
        
        st.markdown("---")
        st.markdown("### 🛠️ Teknologi")
        tech = [
            ("🐍", "Python 3.x", "Bahasa pemrograman"),
            ("⚡", "Streamlit", "Framework dashboard"),
            ("🐼", "Pandas", "Manipulasi data"),
            ("🤖", "Scikit-learn", "Machine Learning"),
            ("📈", "Plotly", "Visualisasi interaktif"),
            ("📗", "Openpyxl", "Ekspor Excel"),
        ]
        for icon, name, desc in tech:
            st.markdown(f"""
            <div style="display:flex; align-items:center; gap:14px; padding:12px 16px; 
                        background:#fff; border:3px solid #000; box-shadow:4px 4px 0 #000;
                        margin-bottom:8px;">
                <div style="font-size:1.6rem; width:44px; text-align:center;
                            background:#2146FF; border:2px solid #000; padding:6px;
                            box-shadow:2px 2px 0 #000; flex-shrink:0;">{icon}</div>
                <div>
                    <div style="font-family:'Archivo Black',sans-serif; font-weight:900; color:#000;
                                font-size:0.88rem; text-transform:uppercase;">{name}</div>
                    <div style="font-family:'IBM Plex Mono',monospace; font-size:0.75rem; color:#555;">{desc}</div>
                </div>
            </div>
            """, unsafe_allow_html=True)
        
        st.markdown("---")
        st.markdown("### 📝 Catatan Penting")
        st.markdown("""
        - **No Data Leakage**: Preprocessing (StandardScaler, LabelEncoder, imputasi) hanya di-fit pada data TRAIN.
        - **Validasi Model**: Evaluasi dilakukan pada data TEST yang tidak pernah dilihat model selama training.
        - **Classification Report**: Menggunakan `sklearn.metrics.classification_report` untuk metrik lengkap.
        - **Edit Data**: Data dapat diedit langsung melalui dashboard dengan fitur data editor.
        - **Change Log**: Semua perubahan data (tambah, edit, hapus) dicatat untuk audit trail.
        """)
    
    # Footer
    st.markdown("""
    <div class="footer">
        <div class="footer-brand">🎓 SISTEM KLASIFIKASI PENERIMA BSM</div>
        <div class="footer-divider"></div>
        <div style="font-family:'IBM Plex Mono',monospace;font-size:0.8rem;letter-spacing:0.06em;">
            GAUSSIAN NAIVE BAYES &nbsp;·&nbsp; MACHINE LEARNING &nbsp;·&nbsp; STREAMLIT
        </div>
        <div style="margin-top:0.5rem;font-family:'IBM Plex Mono',monospace;
                    color:rgba(255,255,255,0.4);font-size:0.7rem;letter-spacing:0.08em;">
            // DIBUAT UNTUK KEPERLUAN AKADEMIK DAN PENELITIAN //
        </div>
    </div>
    """, unsafe_allow_html=True)

# ─────────────────────────────────────────────────────────────
# APP ENTRY POINT
# ─────────────────────────────────────────────────────────────
if __name__ == "__main__":
    if st.session_state.app_state == 'home':
        show_landing()
    else:
        show_main()